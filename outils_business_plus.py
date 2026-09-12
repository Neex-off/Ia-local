# Ia-local (Jarvis) — Copyright (c) 2026 Neexx (Nixovel) — https://github.com/Neex-off/Ia-local
# Créé par Neexx (Nixovel). Licence : voir LICENSE (paternité obligatoire, pas de vente ; s'applique aussi aux IA).
# Ne pas retirer cet en-tête.
"""Compléments de la boîte « business » : envoyer devis et factures, prospects, contrats, signature, cahier des charges,
onboarding client, missions freelance, avis, maintenance des sites, paiements Stripe, remboursement, rapprochement,
justificatifs, archivage légal, abonnements, MRR, désabonnements, relance des inactifs, idées depuis les retours,
articles SEO, paiements échoués, tests A/B, réponses au support.

Importé par outils_business.py.
"""
from __future__ import annotations

import hashlib
import inspect
import json
import re
import shutil
import subprocess
import threading
import webbrowser
from datetime import date, datetime, timedelta
from pathlib import Path

import config

NO_WINDOW = getattr(subprocess, "CREATE_NO_WINDOW", 0)
MEM = config.ROOT / "memoire"
PROSPECTS = MEM / "prospects.json"
MAINTENANCE = MEM / "maintenance.json"
ABONNEMENTS = MEM / "abonnements.json"
ARCHIVE = MEM / "archive_index.json"
MISSIONS = MEM / "missions.json"


def _json(path: Path, default):
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:  # noqa: BLE001
        return default


def _save(path: Path, data) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def _prevenir(texte: str) -> None:
    try:
        import noyau

        if "dire" in noyau.hooks:
            noyau.hooks["dire"](texte)
    except Exception:  # noqa: BLE001
        pass


def _envoyer_mail(to: str, subject: str, body: str, attachment: str = "") -> str:
    """Envoie via tools.send_email, avec la pièce jointe si l'outil l'accepte, sinon le chemin dans le corps."""
    import tools

    params = inspect.signature(tools.send_email).parameters
    if attachment and "attachment" in params:
        return tools.send_email(to, subject, body, attachment=attachment)
    if attachment and "attachments" in params:
        return tools.send_email(to, subject, body, attachments=attachment)
    return tools.send_email(to, subject, body + (f"\n\n(pièce jointe à ajouter : {attachment})" if attachment else ""))


def _stripe(method: str, path: str, **params):
    import requests

    cle = getattr(config, "STRIPE_KEY", "")
    if not cle:
        return None, "Il manque STRIPE_KEY (clé secrète sk_…) dans config.py."
    r = requests.request(method, f"https://api.stripe.com/v1{path}", auth=(cle, ""), params=params if method == "GET" else None,
                         data=params if method != "GET" else None, timeout=30)
    if r.status_code >= 300:
        return None, f"Stripe {r.status_code} : {r.text[:200]}"
    return r.json(), ""


# --------------------------------------------------------------------------
# Devis, factures, contrats
# --------------------------------------------------------------------------

def send_invoice(invoice_number: str, to_email: str, message: str = "") -> str:
    """Envoie un devis ou une facture (PDF créé par create_invoice) par e-mail au client.

    Args:
        invoice_number: Numéro ou morceau du nom du PDF (ex. "F-2026-012").
        to_email: Adresse du client.
        message: Texte d'accompagnement (vide = message standard).
    """
    candidats = [p for d in (config.WORKSPACE, MEM, config.ROOT / "factures", config.WORKSPACE / "factures") if d.exists()
                 for p in d.rglob("*.pdf") if invoice_number.lower() in p.name.lower()]
    if not candidats:
        return f"Aucun PDF contenant « {invoice_number} » : crée-la d'abord avec create_invoice."
    pdf = max(candidats, key=lambda p: p.stat().st_mtime)
    corps = message or f"Bonjour,\n\nVeuillez trouver ci-joint le document {pdf.stem}.\nN'hésitez pas à me contacter pour toute question.\n\nCordialement"
    return _envoyer_mail(to_email, pdf.stem, corps, attachment=str(pdf))


def find_prospects(activity: str, city: str = "", count: int = 20) -> str:
    """Trouve des prospects (commerces, cabinets, artisans…) dans une ville via OpenStreetMap : nom, téléphone, site, adresse ; les garde en mémoire.

    Args:
        activity: Type d'activité (ex. "restaurant", "coiffeur", "plombier", "cabinet d'avocats").
        city: Ville (vide = config ou Mulhouse).
        count: Nombre max.
    """
    import requests

    ville = city or getattr(config, "CITY", "") or "Mulhouse"
    try:
        g = requests.get("https://nominatim.openstreetmap.org/search", params={"q": ville, "format": "json", "limit": 1}, headers={"User-Agent": "jarvis-local"}, timeout=15).json()
        lat, lon = float(g[0]["lat"]), float(g[0]["lon"])
    except Exception as exc:  # noqa: BLE001
        return f"Ville introuvable : {exc}"
    mots = activity.lower()
    tags = {"restaurant": 'amenity="restaurant"', "coiffeur": 'shop="hairdresser"', "plombier": 'craft="plumber"', "avocat": 'office="lawyer"',
            "boulangerie": 'shop="bakery"', "garage": 'shop="car_repair"', "dentiste": 'amenity="dentist"', "kiné": 'healthcare="physiotherapist"',
            "immobilier": 'office="estate_agent"', "salle de sport": 'leisure="fitness_centre"', "hôtel": 'tourism="hotel"', "bar": 'amenity="bar"',
            "pharmacie": 'amenity="pharmacy"', "électricien": 'craft="electrician"', "architecte": 'office="architect"', "comptable": 'office="accountant"'}
    tag = next((v for k, v in tags.items() if k in mots), f'name~"{activity}",i')
    q = f'[out:json][timeout:25];(node[{tag}](around:6000,{lat},{lon});way[{tag}](around:6000,{lat},{lon}););out center {int(count)};'
    try:
        r = requests.post("https://overpass-api.de/api/interpreter", data={"data": q}, timeout=60, headers={"User-Agent": "jarvis-local"}).json()
    except Exception as exc:  # noqa: BLE001
        return f"Overpass injoignable : {exc}"
    base = _json(PROSPECTS, [])
    connus = {p["nom"].lower() for p in base}
    nouveaux = []
    for el in r.get("elements", []):
        t = el.get("tags", {})
        nom = t.get("name")
        if not nom or nom.lower() in connus:
            continue
        p = {"nom": nom, "tel": t.get("phone") or t.get("contact:phone", ""), "site": t.get("website") or t.get("contact:website", ""),
             "email": t.get("email") or t.get("contact:email", ""), "ville": ville, "activite": activity,
             "adresse": " ".join(x for x in (t.get("addr:housenumber", ""), t.get("addr:street", "")) if x), "statut": "nouveau",
             "t": datetime.now().isoformat(timespec="minutes")}
        nouveaux.append(p)
        connus.add(nom.lower())
    base += nouveaux
    _save(PROSPECTS, base)
    if not nouveaux:
        return f"Aucun nouveau prospect « {activity} » à {ville} (ou déjà tous connus : {len(base)} en mémoire)."
    return f"{len(nouveaux)} prospects ajoutés :\n" + "\n".join(
        f"- {p['nom']} — {p['tel'] or 'pas de tél'} — {p['site'] or 'pas de site'}" + (f" — {p['email']}" if p["email"] else "") for p in nouveaux[:count])


def contact_prospects(names: str = "", subject: str = "", template: str = "", send: bool = False) -> str:
    """Contacte des prospects : prépare (ou envoie) un e-mail personnalisé à ceux qui ont une adresse ; les autres sont listés avec leur téléphone.

    Args:
        names: Noms séparés par des virgules (vide = tous les prospects « nouveau »).
        subject: Objet du mail.
        template: Texte avec {nom} et {ville} comme variables.
        send: True pour envoyer vraiment, False pour seulement préparer.
    """
    base = _json(PROSPECTS, [])
    voulus = {n.strip().lower() for n in names.split(",") if n.strip()}
    cibles = [p for p in base if p["nom"].lower() in voulus] if voulus else [p for p in base if p.get("statut") == "nouveau"]
    if not cibles:
        return "Aucun prospect ciblé (find_prospects d'abord)."
    gabarit = template or ("Bonjour,\n\nJe suis développeur à {ville} et j'accompagne des entreprises comme {nom} pour leur site et leurs outils en ligne. "
                           "Auriez-vous 15 minutes cette semaine pour en parler ?\n\nBonne journée")
    faits, sans_mail = [], []
    for p in cibles:
        texte = gabarit.format(nom=p["nom"], ville=p.get("ville", ""))
        if p.get("email"):
            if send:
                faits.append(f"{p['nom']} : {_envoyer_mail(p['email'], subject or 'Prise de contact', texte)[:60]}")
                p["statut"] = "contacté"
                p["contacte_le"] = date.today().isoformat()
            else:
                faits.append(f"{p['nom']} ({p['email']}) : prêt")
        else:
            sans_mail.append(f"{p['nom']} : {p.get('tel') or 'pas de tél'} {p.get('site', '')}")
    if send:
        _save(PROSPECTS, base)
    entete = "Envoyés" if send else "Préparés (send=True pour envoyer)"
    return ((f"{entete} :\n" + "\n".join(faits)) if faits else "Aucun prospect avec e-mail.") + \
        ("\n\nSans e-mail, à appeler ou à contacter via le site :\n" + "\n".join(sans_mail) if sans_mail else "")


def contract_from_template(template: str = "", fields: str = "{}", output: str = "") -> str:
    """Génère un contrat depuis un modèle Word (memoire/modeles/<nom>.docx avec des {{champs}}) ou un modèle intégré de prestation freelance.

    Args:
        template: Nom du modèle dans memoire/modeles (vide = modèle intégré « prestation »).
        fields: JSON des champs, ex. '{"client": "SAS Dupont", "prestation": "site vitrine", "prix": "1 500 € HT", "delai": "4 semaines", "date": "12/09/2026"}'.
        output: Fichier .docx de sortie (vide = workspace/contrat-<client>.docx).
    """
    from docx import Document

    try:
        champs = json.loads(fields or "{}")
    except Exception as exc:  # noqa: BLE001
        return f"JSON invalide : {exc}"
    client = re.sub(r"[^a-z0-9]+", "-", str(champs.get("client", "client")).lower()).strip("-")
    out = Path(output) if output else config.WORKSPACE / f"contrat-{client}.docx"
    out.parent.mkdir(parents=True, exist_ok=True)
    modele = MEM / "modeles" / f"{template}.docx" if template else None
    if modele and modele.exists():
        doc = Document(str(modele))
        manquants = set()
        for p in doc.paragraphs:
            for k, v in champs.items():
                if "{{" + k + "}}" in p.text:
                    for r in p.runs:
                        r.text = r.text.replace("{{" + k + "}}", str(v))
                    if "{{" + k + "}}" in p.text:
                        p.text = p.text.replace("{{" + k + "}}", str(v))
            for m in re.findall(r"{{(\w+)}}", p.text):
                manquants.add(m)
        doc.save(str(out))
        return f"Contrat généré : {out}" + (f" (champs non remplis : {', '.join(manquants)})" if manquants else "")
    prestataire = getattr(config, "USER_NAME", "") or "le Prestataire"
    doc = Document()
    doc.add_heading("Contrat de prestation de services", 0)
    doc.add_paragraph(f"Entre {prestataire} (le Prestataire) et {champs.get('client', '[CLIENT]')} (le Client), il est convenu ce qui suit.")
    sections = [("1. Objet", f"Le Prestataire réalise pour le Client : {champs.get('prestation', '[PRESTATION]')}."),
                ("2. Durée et délais", f"La prestation est livrée sous {champs.get('delai', '[DÉLAI]')} à compter de la signature et du versement de l'acompte."),
                ("3. Prix et paiement", f"Le prix est fixé à {champs.get('prix', '[PRIX]')}. Acompte de {champs.get('acompte', '30 %')} à la commande, solde à la livraison, paiement à 30 jours."),
                ("4. Obligations", "Le Client fournit les contenus et accès nécessaires. Le Prestataire livre un travail conforme aux règles de l'art."),
                ("5. Propriété intellectuelle", "Les livrables sont cédés au Client après paiement intégral. Le Prestataire peut citer la réalisation comme référence."),
                ("6. Révisions", f"{champs.get('revisions', 'Deux')} séries de retours sont incluses ; au-delà, facturation au temps passé."),
                ("7. Résiliation", "Chaque partie peut résilier par écrit avec 15 jours de préavis ; le travail réalisé reste dû."),
                ("8. Droit applicable", "Droit français. Tribunal compétent : celui du siège du Prestataire.")]
    for titre, texte in sections:
        doc.add_heading(titre, 2)
        doc.add_paragraph(texte)
    doc.add_paragraph(f"\nFait le {champs.get('date', date.today().strftime('%d/%m/%Y'))}, en deux exemplaires.\n\nLe Prestataire\t\t\t\tLe Client")
    doc.save(str(out))
    return f"Contrat généré : {out}. Relis-le et adapte les clauses avant envoi (send_contract)."


def send_contract(path: str, to_email: str, message: str = "") -> str:
    """Envoie un contrat à signer : par Yousign (YOUSIGN_KEY) avec signature électronique, sinon par e-mail avec instructions.

    Args:
        path: Fichier du contrat (PDF de préférence).
        to_email: Adresse du signataire.
        message: Message d'accompagnement.
    """
    p = Path(path).expanduser()
    if not p.exists():
        return f"Introuvable : {p}"
    corps_defaut = "Bonjour,\n\nVeuillez trouver ci-joint le contrat. Merci de le signer (signature manuscrite scannée ou signature PDF) et de me le retourner.\n\nCordialement"
    cle = getattr(config, "YOUSIGN_KEY", "")
    if cle and p.suffix.lower() == ".pdf":
        import requests

        base = "https://api.yousign.app/v3"
        h = {"Authorization": f"Bearer {cle}"}
        try:
            sr = requests.post(f"{base}/signature_requests", headers=h, json={"name": p.stem, "delivery_mode": "email"}, timeout=30).json()
            doc = requests.post(f"{base}/signature_requests/{sr['id']}/documents", headers=h, files={"file": (p.name, p.read_bytes(), "application/pdf")},
                                data={"nature": "signable_document"}, timeout=60).json()
            requests.post(f"{base}/signature_requests/{sr['id']}/signers", headers=h, timeout=30,
                          json={"info": {"first_name": "Client", "last_name": to_email.split("@")[0], "email": to_email, "locale": "fr"},
                                "signature_level": "electronic_signature", "signature_authentication_mode": "no_otp",
                                "fields": [{"document_id": doc["id"], "type": "signature", "page": 1, "x": 350, "y": 700}]})
            requests.post(f"{base}/signature_requests/{sr['id']}/activate", headers=h, timeout=30)
            return f"Demande de signature Yousign envoyée à {to_email} ({sr['id']})."
        except Exception as exc:  # noqa: BLE001
            return f"Yousign a échoué ({exc}) : envoi par e-mail à la place. " + _envoyer_mail(to_email, f"Contrat à signer : {p.stem}", message or corps_defaut, attachment=str(p))
    return _envoyer_mail(to_email, f"Contrat à signer : {p.stem}", message or corps_defaut, attachment=str(p)) + " (signature électronique possible avec YOUSIGN_KEY dans config.py)"


def spec_from_call(audio_path: str = "", transcript: str = "", client: str = "") -> str:
    """Cahier des charges depuis un appel : transcrit l'enregistrement (ou prend une transcription) et prépare la structure que tu remplis, écrite en Markdown.

    Args:
        audio_path: Enregistrement de l'appel (transcribe_call produit un WAV).
        transcript: Transcription déjà faite (si pas d'audio).
        client: Nom du client.
    """
    texte = transcript
    if audio_path and not texte:
        try:
            import tools

            texte = tools.transcribe_media(audio_path)
        except Exception as exc:  # noqa: BLE001
            return f"Transcription impossible : {exc}"
    if not texte:
        return "Donne un audio ou une transcription."
    slug = re.sub(r"[^a-z0-9]+", "-", (client or "client").lower()).strip("-")
    out = config.WORKSPACE / f"cahier-des-charges-{slug}.md"
    out.parent.mkdir(parents=True, exist_ok=True)
    gabarit = (f"# Cahier des charges — {client or 'Client'} ({date.today():%d/%m/%Y})\n\n## 1. Contexte et objectifs\n_à rédiger_\n\n## 2. Cibles / utilisateurs\n_à rédiger_\n\n"
               "## 3. Fonctionnalités attendues (priorisées)\n- [ ] …\n\n## 4. Contraintes (techniques, légales, design)\n_à rédiger_\n\n## 5. Livrables\n_à rédiger_\n\n"
               "## 6. Planning et jalons\n_à rédiger_\n\n## 7. Budget\n_à rédiger_\n\n## 8. Points à clarifier\n- …\n\n---\n## Transcription de l'appel\n" + texte[:8000])
    out.write_text(gabarit, encoding="utf-8")
    return (f"Gabarit écrit dans {out} avec la transcription en annexe. Remplis chaque section à partir de l'appel (write_file), en citant les phrases du client "
            "pour les fonctionnalités et en listant les points flous dans « Points à clarifier ».")


def client_onboarding(name: str, email: str, project: str, send_welcome: bool = True, kickoff_days: int = 2) -> str:
    """Onboarding client automatique : fiche client, dossiers du projet, e-mail de bienvenue avec les prochaines étapes, réunion de lancement dans l'agenda.

    Args:
        name: Nom du client.
        email: Son e-mail.
        project: Le projet (une phrase).
        send_welcome: Envoyer l'e-mail de bienvenue.
        kickoff_days: Réunion de lancement dans N jours.
    """
    import tools

    out = []
    try:
        out.append(tools.add_client(name, email)[:80])
    except Exception as exc:  # noqa: BLE001
        out.append(f"fiche client : {exc}")
    slug = re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")
    base = config.WORKSPACE / "clients" / slug
    for d in ("contrat", "devis-factures", "livrables", "echanges", "sources"):
        (base / d).mkdir(parents=True, exist_ok=True)
    (base / "README.md").write_text(f"# {name}\n\nProjet : {project}\nDébut : {date.today():%d/%m/%Y}\nContact : {email}\n", encoding="utf-8")
    out.append(f"dossiers créés dans {base}")
    kick = (date.today() + timedelta(days=kickoff_days)).isoformat()
    try:
        out.append(tools.add_event(f"Lancement projet {name}", kick, "10:00")[:80])
    except Exception:  # noqa: BLE001
        pass
    if send_welcome:
        corps = (f"Bonjour,\n\nMerci pour votre confiance ! Voici les prochaines étapes pour « {project} » :\n"
                 f"1. Réunion de lancement le {datetime.fromisoformat(kick):%d/%m} à 10h (30 min) : objectifs, contenus, accès.\n"
                 "2. Vous m'envoyez logo, textes, images et accès existants.\n3. Je vous livre une première version pour retours.\n\nÀ très vite,")
        out.append(_envoyer_mail(email, f"Bienvenue — lancement de {project}", corps)[:80])
    return "Onboarding fait : " + " ; ".join(out)


def find_missions(keywords: str, platforms: str = "malt,freelance-informatique,codeur,upwork", remote_only: bool = True) -> str:
    """Trouve des missions freelance : ouvre les recherches sur les plateformes et cherche les annonces récentes sur le web ; garde la recherche en mémoire.

    Args:
        keywords: Mots-clés (ex. "développeur python", "site vitrine wordpress").
        platforms: Plateformes séparées par des virgules.
        remote_only: Privilégier le télétravail.
    """
    import urllib.parse

    q = urllib.parse.quote(keywords)
    urls = {"malt": f"https://www.malt.fr/s?q={q}", "freelance-informatique": f"https://www.freelance-informatique.fr/missions-freelance?q={q}",
            "codeur": f"https://www.codeur.com/projects?q={q}", "upwork": f"https://www.upwork.com/nx/search/jobs/?q={q}",
            "comet": "https://app.comet.co/freelancer/missions", "linkedin": f"https://www.linkedin.com/jobs/search/?keywords={q}&f_WT=2"}
    ouverts = []
    for p in platforms.split(","):
        p = p.strip().lower()
        if p in urls:
            webbrowser.open(urls[p])
            ouverts.append(p)
    try:
        import tools

        web = tools.web_search(f"mission freelance {keywords} {'télétravail' if remote_only else ''} {datetime.now():%B %Y}")[:1500]
    except Exception as exc:  # noqa: BLE001
        web = f"recherche impossible : {exc}"
    base = _json(MISSIONS, [])
    base.append({"recherche": keywords, "t": datetime.now().isoformat(timespec="minutes"), "plateformes": ouverts})
    _save(MISSIONS, base[-50:])
    return f"Recherches ouvertes sur {', '.join(ouverts)} : lis les annonces à l'écran (see_screen).\n\nSur le web :\n{web}"


def mission_application(mission_text: str, profile: str = "") -> str:
    """Brouillon de candidature à une mission : analyse l'annonce (besoins, techno, budget, délai) et te fait rédiger une réponse courte et ciblée.

    Args:
        mission_text: Texte de l'annonce.
        profile: Ton profil en 3 lignes (vide = celui en mémoire).
    """
    besoins = re.findall(r"(?i)\b(react|vue|angular|next|node|python|django|fastapi|php|laravel|wordpress|shopify|flutter|react native|swift|kotlin|sql|postgres|mongo|aws|docker|figma|seo|api|ia|llm)\b", mission_text)
    budget = re.search(r"(\d[\d\s]{2,})\s*(€|euros|k€)", mission_text)
    delai = re.search(r"(?i)(urgent|asap|d[ée]s que possible|sous \d+ (jours|semaines)|\d+ (jours|semaines|mois))", mission_text)
    base = _json(MISSIONS, [])
    base.append({"candidature": mission_text[:200], "t": datetime.now().isoformat(timespec="minutes")})
    _save(MISSIONS, base[-50:])
    technos = ", ".join(sorted(set(b.lower() for b in besoins))) or "non précisées"
    return (f"Technos citées : {technos}. Budget : {budget.group(0) if budget else 'non indiqué'}. Délai : {delai.group(0) if delai else 'non indiqué'}.\n"
            + (f"Profil : {profile}\n" if profile else "")
            + "Rédige la candidature (120 à 180 mots) : 1) reformule leur besoin en une phrase pour montrer que tu as compris, 2) une réalisation similaire avec un résultat chiffré, "
            "3) comment tu procéderais (3 étapes), 4) disponibilité et tarif ou fourchette, 5) une question précise sur le projet. Pas de flatterie, pas de « je suis passionné ».")


def review_request(client_name: str, email: str, link: str = "", project: str = "") -> str:
    """Demande d'avis après livraison : e-mail court et chaleureux avec le lien (Google, Malt, LinkedIn…).

    Args:
        client_name: Nom du client.
        email: Son adresse.
        link: Lien où laisser l'avis (vide = réponse par mail).
        project: Le projet livré.
    """
    corps = (f"Bonjour {client_name},\n\nJ'espère que {project or 'le projet'} vous donne entière satisfaction. Si vous avez deux minutes, un avis m'aiderait énormément"
             + (f" : {link}" if link else " (une réponse à ce mail suffit)") + ".\n\nMerci encore pour votre confiance, et à bientôt.\n")
    return _envoyer_mail(email, "Un petit avis ?", corps)


def maintenance_reminders(action: str = "check", site: str = "", every_days: int = 30) -> str:
    """Rappels de maintenance des sites clients (mises à jour, sauvegardes, certificats) : ajouter, lister, vérifier ce qui est dû, marquer fait.

    Args:
        action: "add", "list", "check", "done" ou "remove".
        site: Nom ou domaine du site.
        every_days: Fréquence en jours (add).
    """
    base = _json(MAINTENANCE, [])
    auj = date.today()
    if action == "add":
        base = [m for m in base if m["site"] != site]
        base.append({"site": site, "tous_les_jours": int(every_days), "prochaine": (auj + timedelta(days=every_days)).isoformat()})
        _save(MAINTENANCE, base)
        return f"Maintenance de {site} tous les {every_days} jours, prochaine le {base[-1]['prochaine']}."
    if action == "remove":
        base = [m for m in base if m["site"] != site]
        _save(MAINTENANCE, base)
        return f"{site} retiré."
    if action == "done":
        for m in base:
            if m["site"] == site:
                m["prochaine"] = (auj + timedelta(days=m["tous_les_jours"])).isoformat()
                m["derniere"] = auj.isoformat()
        _save(MAINTENANCE, base)
        return f"Maintenance de {site} notée faite."
    if action == "list":
        return "\n".join(f"- {m['site']} : tous les {m['tous_les_jours']} j, prochaine {m['prochaine']}" for m in base) or "Aucun site suivi."
    dues = [m for m in base if m["prochaine"] <= (auj + timedelta(days=3)).isoformat()]
    if not dues:
        return "Aucune maintenance due dans les 3 jours."
    lignes = []
    for m in dues:
        try:
            import tools

            etat = tools.site_uptime(m["site"])[:80]
        except Exception:  # noqa: BLE001
            etat = ""
        lignes.append(f"- {m['site']} (prévue {m['prochaine']}) {etat}")
    return "Maintenances à faire :\n" + "\n".join(lignes) + "\nCheck-list : sauvegarde, mises à jour CMS/plugins, certificat, formulaire de contact, vitesse."


# --------------------------------------------------------------------------
# Paiements et compta
# --------------------------------------------------------------------------

def stripe_payments(action: str = "list", days: int = 30, customer: str = "") -> str:
    """Suivi des paiements Stripe : derniers paiements, solde, ou recherche d'un client.

    Args:
        action: "list", "balance" ou "customer".
        days: Fenêtre pour list.
        customer: E-mail ou nom pour customer.
    """
    if action == "balance":
        d, err = _stripe("GET", "/balance")
        if err:
            return err
        dispo = ", ".join(f"{b['amount']/100:.2f} {b['currency'].upper()} disponible" for b in d.get("available", []))
        attente = ", ".join(f"{b['amount']/100:.2f} {b['currency'].upper()} en attente" for b in d.get("pending", []))
        return f"Solde : {dispo} ; {attente}"
    if action == "customer":
        d, err = _stripe("GET", "/customers/search", query=f"email:'{customer}'" if "@" in customer else f"name~'{customer}'")
        if err:
            return err
        return "\n".join(f"- {c.get('name')} {c.get('email')} ({c['id']})" for c in d.get("data", [])) or "Aucun client."
    depuis = int((datetime.now() - timedelta(days=days)).timestamp())
    d, err = _stripe("GET", "/payment_intents", limit=50, **{"created[gte]": depuis})
    if err:
        return err
    lignes = []
    total = 0.0
    for pi in d.get("data", []):
        if pi.get("status") == "succeeded":
            total += pi["amount"] / 100
        lignes.append(f"- {datetime.fromtimestamp(pi['created']):%d/%m} {pi['amount']/100:.2f} {pi['currency'].upper()} {pi['status']} {pi.get('description') or ''} ({pi['id']})")
    return f"{len(lignes)} paiements sur {days} jours, {total:.2f} encaissés :\n" + "\n".join(lignes[:30])


def refund_payment(payment_id: str, amount_eur: float = 0, reason: str = "requested_by_customer") -> str:
    """Rembourse un paiement Stripe (total, ou partiel avec un montant).

    Args:
        payment_id: Id du PaymentIntent (pi_…) ou de la charge (ch_…).
        amount_eur: Montant à rembourser (0 = tout).
        reason: duplicate, fraudulent ou requested_by_customer.
    """
    params = {"reason": reason}
    params["payment_intent" if payment_id.startswith("pi_") else "charge"] = payment_id
    if amount_eur:
        params["amount"] = int(round(float(amount_eur) * 100))
    d, err = _stripe("POST", "/refunds", **params)
    if err:
        return err
    return f"Remboursement {d['id']} : {d['amount']/100:.2f} {d['currency'].upper()}, statut {d['status']}."


def reconcile_payments(csv_path: str = "", days: int = 60) -> str:
    """Rapprochement factures ↔ paiements : compare tes factures impayées avec les paiements Stripe (ou un relevé CSV date;montant;libellé) et propose les correspondances.

    Args:
        csv_path: Relevé bancaire CSV (vide = Stripe).
        days: Fenêtre de recherche.
    """
    import tools

    try:
        impayees = tools.unpaid_report()
    except Exception as exc:  # noqa: BLE001
        return f"Impossible de lire les factures : {exc}"
    factures = [(m.group(1), float(m.group(2).replace(" ", "").replace(",", "."))) for m in re.finditer(r"(\S+)\D{0,40}?(\d[\d\s]*[.,]?\d*)\s*€", impayees)]
    if not factures:
        return "Aucune facture impayée à rapprocher (ou format non lu) :\n" + impayees[:500]
    paiements = []
    if csv_path:
        import csv

        p = Path(csv_path).expanduser()
        texte = p.read_text(encoding="utf-8", errors="replace")
        sep = ";" if ";" in texte[:500] else ","
        for row in csv.reader(texte.splitlines(), delimiter=sep):
            if len(row) < 2:
                continue
            try:
                montant = float(re.sub(r"[^\d,.-]", "", row[1]).replace(",", "."))
            except ValueError:
                continue
            paiements.append((row[0], montant, row[2] if len(row) > 2 else ""))
    else:
        depuis = int((datetime.now() - timedelta(days=days)).timestamp())
        d, err = _stripe("GET", "/payment_intents", limit=100, **{"created[gte]": depuis})
        if err:
            return err
        paiements = [(datetime.fromtimestamp(pi["created"]).strftime("%d/%m"), pi["amount"] / 100, pi.get("description") or pi["id"]) for pi in d.get("data", []) if pi.get("status") == "succeeded"]
    out = []
    for num, montant in factures:
        m = [p for p in paiements if abs(p[1] - montant) < 0.01]
        if m:
            out.append(f"- facture {num} ({montant:.2f} €) ↔ paiement du {m[0][0]} « {m[0][2]} » : mark_paid(\"{num}\") ?")
        else:
            out.append(f"- facture {num} ({montant:.2f} €) : aucun paiement de ce montant")
    return "Rapprochement :\n" + "\n".join(out)


def receipts_by_month(folder: str, destination: str = "", dry_run: bool = True) -> str:
    """Classe les justificatifs (factures, reçus, tickets) par mois : dossiers AAAA-MM d'après la date du document (nom, contenu PDF, ou date du fichier).

    Args:
        folder: Dossier des justificatifs en vrac.
        destination: Dossier de classement (vide = le même).
        dry_run: True = plan, False = déplacer.
    """
    p = Path(folder).expanduser()
    if not p.is_dir():
        return f"Dossier introuvable : {p}"
    dest = Path(destination).expanduser() if destination else p
    plan: dict[str, list[Path]] = {}
    for f in p.iterdir():
        if not f.is_file() or f.suffix.lower() not in (".pdf", ".jpg", ".jpeg", ".png", ".xlsx", ".csv", ".txt"):
            continue
        mois = None
        m = re.search(r"(20\d{2})[-_ ]?(0[1-9]|1[0-2])", f.name)
        if m:
            mois = f"{m.group(1)}-{m.group(2)}"
        elif f.suffix.lower() == ".pdf":
            try:
                from pypdf import PdfReader

                txt = (PdfReader(str(f)).pages[0].extract_text() or "")[:3000]
                m2 = re.search(r"(0[1-9]|[12]\d|3[01])[/.](0[1-9]|1[0-2])[/.](20\d{2})", txt)
                if m2:
                    mois = f"{m2.group(3)}-{m2.group(2)}"
            except Exception:  # noqa: BLE001
                pass
        mois = mois or datetime.fromtimestamp(f.stat().st_mtime).strftime("%Y-%m")
        plan.setdefault(mois, []).append(f)
    if not plan:
        return "Aucun justificatif."
    if dry_run:
        return "\n".join(f"- {k} : {len(v)} fichier(s)" for k, v in sorted(plan.items())) + "\nRappelle avec dry_run=False pour classer."
    n = 0
    for k, fs in plan.items():
        d = dest / k
        d.mkdir(parents=True, exist_ok=True)
        for f in fs:
            shutil.move(str(f), str(d / f.name))
            n += 1
    return f"{n} justificatifs classés dans {len(plan)} mois sous {dest}."


def legal_archive(folder: str, archive_dir: str = "", years: int = 10) -> str:
    """Archivage légal des factures : copie dans un dossier d'archives par année, empreinte SHA-256 de chaque fichier dans un index, fichiers passés en lecture seule.

    Args:
        folder: Dossier des factures à archiver.
        archive_dir: Dossier d'archives (vide = Documents/Archives-factures).
        years: Durée de conservation (information).
    """
    import os
    import stat

    p = Path(folder).expanduser()
    if not p.is_dir():
        return f"Dossier introuvable : {p}"
    arch = Path(archive_dir).expanduser() if archive_dir else Path.home() / "Documents" / "Archives-factures"
    index = _json(ARCHIVE, {})
    n, deja = 0, 0
    for f in p.rglob("*"):
        if not f.is_file() or f.suffix.lower() not in (".pdf", ".xml", ".jpg", ".png"):
            continue
        sha = hashlib.sha256(f.read_bytes()).hexdigest()
        if sha in index:
            deja += 1
            continue
        m = re.search(r"(20\d{2})", f.name)
        annee = m.group(1) if m else datetime.fromtimestamp(f.stat().st_mtime).strftime("%Y")
        d = arch / annee
        d.mkdir(parents=True, exist_ok=True)
        cible = d / f.name
        if cible.exists():
            cible = d / f"{f.stem}-{sha[:8]}{f.suffix}"
        shutil.copy2(f, cible)
        os.chmod(cible, stat.S_IREAD)
        index[sha] = {"fichier": str(cible), "source": str(f), "t": datetime.now().isoformat(timespec="minutes")}
        n += 1
    _save(ARCHIVE, index)
    return f"{n} fichier(s) archivés dans {arch} ({deja} déjà présents), index d'empreintes à jour ({len(index)} au total). Conservation : {years} ans."


def subscription_reminders(action: str = "check", name: str = "", amount: float = 0, renew_date: str = "", period: str = "mensuel", days: int = 7) -> str:
    """Abonnements (logiciels, hébergement, domaines, salle…) : ajouter, lister, coût mensuel total, et rappels avant renouvellement.

    Args:
        action: "add", "remove", "list" ou "check".
        name: Nom de l'abonnement.
        amount: Montant.
        renew_date: Prochaine date de renouvellement AAAA-MM-JJ.
        period: "mensuel" ou "annuel".
        days: Fenêtre d'alerte (check).
    """
    base = _json(ABONNEMENTS, [])
    if action == "add":
        base = [a for a in base if a["nom"].lower() != name.lower()]
        base.append({"nom": name, "montant": float(amount), "periode": period, "renouvellement": renew_date})
        _save(ABONNEMENTS, base)
        return f"Abonnement « {name} » : {amount} € {period}, renouvellement le {renew_date}."
    if action == "remove":
        base = [a for a in base if a["nom"].lower() != name.lower()]
        _save(ABONNEMENTS, base)
        return f"« {name} » retiré."
    auj = date.today()
    change = False
    for a in base:
        try:
            d = date.fromisoformat(a["renouvellement"])
            while d < auj:
                d = d.replace(year=d.year + 1) if a["periode"] == "annuel" else d + timedelta(days=30)
                a["renouvellement"] = d.isoformat()
                change = True
        except Exception:  # noqa: BLE001
            continue
    if change:
        _save(ABONNEMENTS, base)
    mensuel = sum(a["montant"] if a["periode"] == "mensuel" else a["montant"] / 12 for a in base)
    if action == "list":
        lignes = "\n".join(f"- {a['nom']} : {a['montant']} € {a['periode']}, prochain {a['renouvellement']}" for a in sorted(base, key=lambda x: x["renouvellement"]))
        return (lignes or "Aucun abonnement.") + f"\nTotal : {mensuel:.2f} € par mois ({mensuel*12:.0f} € par an)."
    proches = [a for a in base if a["renouvellement"] <= (auj + timedelta(days=days)).isoformat()]
    if not proches:
        return f"Aucun renouvellement dans les {days} jours ({mensuel:.2f} € par mois au total)."
    return "Renouvellements proches :\n" + "\n".join(f"- {a['nom']} : {a['montant']} € le {a['renouvellement']}" for a in proches) + "\nÀ garder ? Sinon résilie avant la date."


def mrr_report() -> str:
    """Suivi du chiffre d'affaires récurrent (MRR) : abonnements actifs Stripe par plan, MRR total, nouveaux du mois."""
    d, err = _stripe("GET", "/subscriptions", status="active", limit=100, **{"expand[]": "data.items.data.price"})
    if err:
        return err
    mrr = 0.0
    par_plan: dict[str, float] = {}
    nouveaux = 0
    debut_mois = datetime.now().replace(day=1).timestamp()
    for s in d.get("data", []):
        for it in s["items"]["data"]:
            pr = it["price"]
            montant = (pr.get("unit_amount") or 0) / 100 * it.get("quantity", 1)
            if pr.get("recurring", {}).get("interval") == "year":
                montant /= 12
            mrr += montant
            nom = pr.get("nickname") or pr["id"]
            par_plan[nom] = par_plan.get(nom, 0) + montant
        if s["created"] >= debut_mois:
            nouveaux += 1
    return (f"MRR : {mrr:.2f} € ({len(d.get('data', []))} abonnements actifs, {nouveaux} nouveaux ce mois, ARR {mrr*12:.0f} €).\n"
            + "\n".join(f"- {k} : {v:.2f} €/mois" for k, v in sorted(par_plan.items(), key=lambda x: -x[1])))


def churn_report(days: int = 30) -> str:
    """Suivi des désabonnements : abonnements Stripe annulés sur la période, taux de churn, raisons si renseignées.

    Args:
        days: Période.
    """
    depuis = int((datetime.now() - timedelta(days=days)).timestamp())
    annules, err = _stripe("GET", "/subscriptions", status="canceled", limit=100, **{"created[gte]": int((datetime.now() - timedelta(days=365)).timestamp())})
    if err:
        return err
    actifs, _ = _stripe("GET", "/subscriptions", status="active", limit=100)
    recents = [s for s in annules.get("data", []) if (s.get("canceled_at") or 0) >= depuis]
    n_actifs = len((actifs or {}).get("data", []))
    taux = len(recents) / max(1, n_actifs + len(recents)) * 100
    raisons: dict[str, int] = {}
    for s in recents:
        det = s.get("cancellation_details") or {}
        r = det.get("feedback") or det.get("reason") or "non précisée"
        raisons[r] = raisons.get(r, 0) + 1
    return (f"{len(recents)} désabonnement(s) sur {days} jours, {n_actifs} actifs : churn {taux:.1f} %.\nRaisons : "
            + (", ".join(f"{k} ({v})" for k, v in raisons.items()) or "aucune"))


def inactive_users(days: int = 30, dsn: str = "", table: str = "users", email_col: str = "email", last_col: str = "last_login", send: bool = False, subject: str = "", body: str = "") -> str:
    """Relance des utilisateurs inactifs : les trouve dans ta base (PostgreSQL) et, si demandé, leur envoie un e-mail de relance.

    Args:
        days: Inactifs depuis N jours.
        dsn: Chaîne de connexion (vide = config.DB_DSN).
        table: Table des utilisateurs.
        email_col: Colonne e-mail.
        last_col: Colonne de dernière activité.
        send: True pour envoyer les mails.
        subject: Objet du mail.
        body: Corps (avec {email} possible).
    """
    try:
        from outils_code_plus import _connect

        with _connect(dsn) as conn, conn.cursor() as cur:
            cur.execute(f"SELECT {email_col}, {last_col} FROM {table} WHERE {last_col} < NOW() - INTERVAL '{int(days)} days' ORDER BY {last_col} LIMIT 200")
            rows = cur.fetchall()
    except Exception as exc:  # noqa: BLE001
        return f"Base injoignable : {exc}"
    if not rows:
        return f"Aucun utilisateur inactif depuis {days} jours."
    if not send:
        return (f"{len(rows)} inactifs depuis {days} jours :\n" + "\n".join(f"- {r[0]} (dernière activité {str(r[1])[:10]})" for r in rows[:40])
                + "\nRappelle avec send=True, subject et body pour les relancer.")
    n = 0
    texte = body or "Bonjour,\n\nVotre compte vous attend : de nouvelles fonctionnalités sont arrivées. Revenez quand vous voulez !\n"
    for r in rows:
        try:
            _envoyer_mail(r[0], subject or "On ne vous a pas vu depuis un moment", texte.format(email=r[0]))
            n += 1
        except Exception:  # noqa: BLE001
            pass
    return f"{n} relance(s) envoyée(s)."


def feature_ideas(feedback: str = "", path: str = "") -> str:
    """Idées de fonctionnalités depuis les retours utilisateurs (texte ou fichier) : regroupe par thème, compte les demandes, fait ressortir les plus fréquentes.

    Args:
        feedback: Retours collés (un par ligne).
        path: Ou un fichier texte / CSV de retours.
    """
    if path:
        p = Path(path).expanduser()
        if not p.exists():
            return f"Introuvable : {p}"
        feedback = p.read_text(encoding="utf-8", errors="replace")
    lignes = [l.strip() for l in feedback.splitlines() if len(l.strip()) > 10]
    if not lignes:
        return "Donne des retours."
    vides = {"le", "la", "les", "de", "des", "un", "une", "et", "à", "que", "qui", "pour", "pas", "ne", "je", "il", "on", "ce", "c'est", "est", "en", "du", "au",
             "the", "to", "and", "of", "is", "it", "vous", "nous", "avec", "sur", "dans", "plus", "très", "trop", "ça", "cette", "mais", "ou", "si", "serait", "bien", "faire", "peut", "avoir"}
    themes: dict[str, list[int]] = {}
    for i, l in enumerate(lignes):
        mots = [m for m in re.findall(r"[a-zàâçéèêëîïôûùüÿ']{4,}", l.lower()) if m not in vides]
        for j in range(len(mots) - 1):
            themes.setdefault(f"{mots[j]} {mots[j + 1]}", []).append(i)
        for m in mots:
            themes.setdefault(m, []).append(i)
    top = sorted(((len(set(v)), k) for k, v in themes.items() if len(set(v)) >= 2), reverse=True)[:15]
    demandes = [l for l in lignes if re.search(r"(?i)(pourriez|pourrait|serait bien|manque|il faudrait|ajouter|j'aimerais|ce serait|feature|fonctionnalit)", l)]
    return (f"{len(lignes)} retours. Thèmes récurrents : " + ", ".join(f"{k} ({n})" for n, k in top) + "\n\nDemandes explicites :\n"
            + "\n".join(f"- {d[:140]}" for d in demandes[:20]) + "\n\nDéduis 5 à 8 fonctionnalités, chacune avec : nom, problème résolu, nombre de retours qui la demandent, effort estimé (S/M/L).")


def seo_article_draft(topic: str, keywords: str = "", length: int = 1200) -> str:
    """Brouillon d'article SEO : intention de recherche d'après les premiers résultats, plan H2/H3, mots-clés, méta ; écrit le squelette dans workspace.

    Args:
        topic: Sujet / requête cible.
        keywords: Mots-clés secondaires séparés par des virgules.
        length: Longueur visée en mots.
    """
    serp = ""
    try:
        import tools

        serp = tools.web_search(topic)[:1500]
    except Exception:  # noqa: BLE001
        pass
    slug = re.sub(r"[^a-z0-9]+", "-", topic.lower()).strip("-")
    out = config.WORKSPACE / f"article-{slug}.md"
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(f"---\ntitle: \"{topic}\"\ndescription: \"_méta description 150 caractères_\"\nkeywords: [{keywords}]\n---\n\n# {topic}\n\n_intro 3 lignes avec le mot-clé_\n\n"
                   "## H2 1\n\n## H2 2\n\n### H3\n\n## FAQ\n\n**Question ?**\n\n## Conclusion\n", encoding="utf-8")
    return (f"Résultats actuels pour « {topic} » :\n{serp}\n\nSquelette écrit dans {out}. Rédige l'article ({length} mots) : titre avec le mot-clé, intro qui répond direct, "
            f"H2 qui reprennent les questions des résultats, mots-clés secondaires ({keywords or 'à choisir'}) placés naturellement, FAQ de 3 questions, méta description. "
            "Écris-le avec write_file dans ce fichier.")


_ECHECS: dict = {"thread": None, "stop": threading.Event()}


def failed_payment_alerts(action: str = "check", days: int = 7) -> str:
    """Alerte paiement échoué (Stripe) : liste les échecs récents, ou surveille en fond et prévient à la voix.

    Args:
        action: "check", "start" ou "stop".
        days: Fenêtre pour check.
    """
    if action == "stop":
        _ECHECS["stop"].set()
        return "Surveillance des paiements arrêtée."
    if action == "start":
        if _ECHECS["thread"] and _ECHECS["thread"].is_alive():
            return "Déjà active."
        _ECHECS["stop"].clear()

        def boucle():
            vus: set = set()
            while not _ECHECS["stop"].is_set():
                d, _ = _stripe("GET", "/events", type="invoice.payment_failed", limit=10)
                for e in (d or {}).get("data", []):
                    if e["id"] not in vus:
                        vus.add(e["id"])
                        inv = e["data"]["object"]
                        _prevenir(f"Paiement échoué : {inv.get('customer_email') or inv.get('customer')} pour {inv.get('amount_due', 0)/100:.2f} euros.")
                _ECHECS["stop"].wait(3600)

        _ECHECS["thread"] = threading.Thread(target=boucle, daemon=True)
        _ECHECS["thread"].start()
        return "Surveillance des paiements échoués lancée (vérification toutes les heures)."
    depuis = int((datetime.now() - timedelta(days=days)).timestamp())
    d, err = _stripe("GET", "/events", type="invoice.payment_failed", limit=50, **{"created[gte]": depuis})
    if err:
        return err
    evts = d.get("data", [])
    if not evts:
        return f"Aucun paiement échoué sur {days} jours."
    lignes = []
    for e in evts[:20]:
        o = e["data"]["object"]
        detail = (o.get("last_finalization_error") or {}).get("message", "sans détail")
        lignes.append(f"- {datetime.fromtimestamp(e['created']):%d/%m} {o.get('customer_email') or o.get('customer')} : {o.get('amount_due', 0)/100:.2f} € ({detail})")
    return f"{len(evts)} paiement(s) échoué(s) :\n" + "\n".join(lignes)


def ab_test(page_path: str, variant_b_path: str = "", split: int = 50, beacon_url: str = "") -> str:
    """Test A/B d'une page statique : crée la variante B (copie à modifier) et un script qui répartit les visiteurs, garde leur variante, et envoie les événements à un endpoint (ou les compte en local).

    Args:
        page_path: Fichier HTML de la page A.
        variant_b_path: Fichier de la variante B (vide = copie « -b.html » à modifier).
        split: Pourcentage de visiteurs envoyés sur B.
        beacon_url: URL qui reçoit les événements (vide = comptage dans localStorage).
    """
    p = Path(page_path).expanduser()
    if not p.exists():
        return f"Introuvable : {p}"
    b = Path(variant_b_path).expanduser() if variant_b_path else p.with_name(p.stem + "-b" + p.suffix)
    if not b.exists():
        shutil.copy(p, b)
    envoi = (f"navigator.sendBeacon('{beacon_url}',JSON.stringify(d));" if beacon_url
             else "var s=JSON.parse(localStorage.getItem(k+'_events')||'[]');s.push(d);localStorage.setItem(k+'_events',JSON.stringify(s));")
    script = (f"<script>(function(){{var k='ab_{p.stem}';var v=localStorage.getItem(k);if(!v){{v=Math.random()*100<{int(split)}?'B':'A';localStorage.setItem(k,v);}}"
              f"if(v==='B'&&!location.pathname.endsWith('{b.name}')){{location.replace('{b.name}'+location.search);return;}}"
              "window.abTrack=function(ev){var d={variant:v,event:ev,page:location.pathname,t:Date.now()};" + envoi + "};abTrack('view');})();</script>")
    for f in (p, b):
        txt = f.read_text(encoding="utf-8", errors="replace")
        if "abTrack" not in txt:
            txt = txt.replace("<head>", "<head>" + script, 1) if "<head>" in txt else script + txt
            f.write_text(txt, encoding="utf-8")
    ou = f"envoyés à {beacon_url}" if beacon_url else f"dans localStorage (clé ab_{p.stem}_events) ou via ton outil d'analytics"
    return (f"Test A/B en place : A = {p.name}, B = {b.name} ({split} % vers B). Modifie {b.name} (titre, bouton, prix…). Appelle abTrack('clic') sur l'action à mesurer "
            f"(ex. onclick=\"abTrack('achat')\"). Résultats : {ou}.")


def support_reply_drafts(messages: str, product: str = "", tone: str = "clair et rassurant") -> str:
    """Brouillons de réponses au support (un par message) que tu rédiges ensuite : accusé, solution ou étapes, délai, clôture.

    Args:
        messages: Les messages reçus, un par ligne ou séparés par ---.
        product: Nom du produit (pour le contexte).
        tone: Ton voulu.
    """
    lignes = [l.strip() for l in re.split(r"\n---\n|\n(?=\S)", messages) if l.strip()]
    if not lignes:
        return "Donne les messages."
    return (f"{len(lignes)} message(s) support" + (f" pour {product}" if product else "") + f". Pour chacun, une réponse {tone} : 1) reformule le problème, 2) solution ou étapes numérotées "
            "(ou question précise si l'info manque), 3) délai réaliste si action de notre côté, 4) phrase de clôture. Pas de jargon, pas d'excuses en boucle. Messages :\n"
            + "\n".join(f"{i}. {l[:400]}" for i, l in enumerate(lignes, start=1)))


TOOLS = [send_invoice, find_prospects, contact_prospects, contract_from_template, send_contract, spec_from_call, client_onboarding,
         find_missions, mission_application, review_request, maintenance_reminders, stripe_payments, refund_payment,
         reconcile_payments, receipts_by_month, legal_archive, subscription_reminders, mrr_report, churn_report, inactive_users,
         feature_ideas, seo_article_draft, failed_payment_alerts, ab_test, support_reply_drafts]
