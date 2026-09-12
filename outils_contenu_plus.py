# Ia-local (Jarvis) — Copyright (c) 2026 Neexx (Nixovel) — https://github.com/Neex-off/Ia-local
# Créé par Neexx (Nixovel). Licence : voir LICENSE (paternité obligatoire, pas de vente ; s'applique aussi aux IA).
# Ne pas retirer cet en-tête.
"""Compléments de la boîte « contenu » : charte graphique, déclinaisons de logo, cartes de visite, Blender, Figma/Canva,
création de contenu (idées, miniatures, stats, publication programmée, formats, commentaires, voix off, tendances,
réponses, meilleure heure, rétention, hashtags), vidéo et audio (recadrage visage, musique libre, sauvegarde des rushs,
égaliseur casque, sons d'ambiance, identification d'une chanson, radio et podcasts).

Importé par outils_contenu.py.
"""
from __future__ import annotations

import json
import re
import shutil
import subprocess
import sys
import threading
import time
import webbrowser
from datetime import datetime
from pathlib import Path

import config

NO_WINDOW = getattr(subprocess, "CREATE_NO_WINDOW", 0)
IS_WINDOWS = sys.platform.startswith("win")
MEM = config.ROOT / "memoire"
MARQUE = MEM / "marque.json"
CONTENU = MEM / "contenu.json"
STATS = MEM / "stats_plateformes.json"


def _json(path: Path, default):
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:  # noqa: BLE001
        return default


def _save(path: Path, data) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def _ffmpeg():
    exe = shutil.which("ffmpeg")
    if exe:
        return exe
    try:
        import imageio_ffmpeg

        return imageio_ffmpeg.get_ffmpeg_exe()
    except Exception:  # noqa: BLE001
        return None


def _run(args: list[str], timeout: float = 600) -> str:
    r = subprocess.run(args, capture_output=True, text=True, timeout=timeout, encoding="utf-8", errors="replace", creationflags=NO_WINDOW)
    return ((r.stdout or "") + (r.stderr or "")).strip()[-1500:]


def _hex(c: str) -> tuple[int, int, int]:
    c = c.lstrip("#")
    return int(c[0:2], 16), int(c[2:4], 16), int(c[4:6], 16)


def _police(taille: int):
    from PIL import ImageFont

    for f in (r"C:\Windows\Fonts\arialbd.ttf", r"C:\Windows\Fonts\segoeuib.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", "/Library/Fonts/Arial Bold.ttf"):
        if Path(f).exists():
            return ImageFont.truetype(f, taille)
    return ImageFont.load_default()


# --------------------------------------------------------------------------
# Design et branding
# --------------------------------------------------------------------------

def brand_kit(action: str = "show", primary: str = "", secondary: str = "", font: str = "", logo: str = "", name: str = "") -> str:
    """Charte graphique : enregistre couleurs, police, logo et nom (set) ou les affiche (show). Utilisée par brand_document, business_cards, thumbnails.

    Args:
        action: "set" ou "show".
        primary: Couleur principale hex (ex. "#1E3A8A").
        secondary: Couleur secondaire hex.
        font: Nom de la police.
        logo: Chemin du logo (PNG de préférence).
        name: Nom de la marque.
    """
    base = _json(MARQUE, {})
    if action == "set":
        for k, v in (("primaire", primary), ("secondaire", secondary), ("police", font), ("logo", logo), ("nom", name)):
            if v:
                base[k] = v
        _save(MARQUE, base)
    if not base:
        return "Aucune charte : brand_kit(\"set\", primary=\"#1E3A8A\", secondary=\"#F59E0B\", font=\"Inter\", logo=\"chemin.png\", name=\"Ma marque\")."
    return "Charte : " + ", ".join(f"{k} = {v}" for k, v in base.items())


def brand_document(kind: str, title: str, content: str, output: str = "") -> str:
    """Crée un document (docx ou pdf) aux couleurs de la charte : titre dans la couleur principale, logo en en-tête, police de la marque.

    Args:
        kind: "docx" ou "pdf".
        title: Titre du document.
        content: Corps du texte (paragraphes séparés par des lignes vides ; lignes commençant par # = sous-titres).
        output: Fichier de sortie (vide = workspace/<titre>.<kind>).
    """
    m = _json(MARQUE, {})
    prim = m.get("primaire", "#1E3A8A")
    police = m.get("police", "Calibri")
    logo = m.get("logo", "")
    slug = re.sub(r"[^a-z0-9]+", "-", title.lower()).strip("-") or "document"
    out = Path(output) if output else config.WORKSPACE / f"{slug}.{kind}"
    out.parent.mkdir(parents=True, exist_ok=True)
    blocs = [b.strip() for b in content.split("\n\n") if b.strip()]
    if kind == "docx":
        from docx import Document
        from docx.shared import Pt, RGBColor

        doc = Document()
        if logo and Path(logo).exists():
            doc.add_picture(logo, width=Pt(90))
        h = doc.add_heading(title, 0)
        for r in h.runs:
            r.font.color.rgb = RGBColor(*_hex(prim))
            r.font.name = police
        for b in blocs:
            if b.startswith("#"):
                hh = doc.add_heading(b.lstrip("# "), 1)
                for r in hh.runs:
                    r.font.color.rgb = RGBColor(*_hex(prim))
            else:
                p = doc.add_paragraph(b)
                for r in p.runs:
                    r.font.name = police
                    r.font.size = Pt(11)
        doc.save(str(out))
    else:
        from fpdf import FPDF

        pdf = FPDF()
        pdf.set_auto_page_break(True, 18)
        pdf.add_page()
        if logo and Path(logo).exists():
            pdf.image(logo, 10, 8, 30)
            pdf.ln(22)
        pdf.set_text_color(*_hex(prim))
        pdf.set_font("Helvetica", "B", 20)
        pdf.multi_cell(0, 10, title)
        pdf.ln(2)
        for b in blocs:
            if b.startswith("#"):
                pdf.set_text_color(*_hex(prim))
                pdf.set_font("Helvetica", "B", 14)
                pdf.multi_cell(0, 8, b.lstrip("# "))
            else:
                pdf.set_text_color(30, 30, 30)
                pdf.set_font("Helvetica", "", 11)
                pdf.multi_cell(0, 6, b)
            pdf.ln(2)
        pdf.output(str(out))
    return f"Document créé : {out}"


def logo_variants(logo_path: str, brand_name: str = "", output_dir: str = "") -> str:
    """Décline un logo : versions blanche, noire, monochrome couleur de marque, icône carrée, horizontale avec le nom, favicon.

    Args:
        logo_path: Logo source (PNG avec transparence de préférence).
        brand_name: Nom pour la version horizontale (vide = charte).
        output_dir: Dossier de sortie (vide = workspace/logo).
    """
    from PIL import Image, ImageDraw

    p = Path(logo_path).expanduser()
    if not p.exists():
        return f"Introuvable : {p}"
    m = _json(MARQUE, {})
    nom = brand_name or m.get("nom", "")
    out = Path(output_dir) if output_dir else config.WORKSPACE / "logo"
    out.mkdir(parents=True, exist_ok=True)
    img = Image.open(p).convert("RGBA")
    if img.getextrema()[3][0] == 255:  # pas de transparence : on retire la couleur du coin
        fond = img.getpixel((0, 0))
        img.putdata([(0, 0, 0, 0) if all(abs(px[i] - fond[i]) < 25 for i in range(3)) else px for px in img.getdata()])
    alpha = img.split()[3]
    faits = []

    def teinte(couleur, nomf):
        solid = Image.new("RGBA", img.size, couleur + (255,))
        solid.putalpha(alpha)
        solid.save(out / nomf)
        faits.append(nomf)

    teinte((255, 255, 255), "logo-blanc.png")
    teinte((0, 0, 0), "logo-noir.png")
    if m.get("primaire"):
        teinte(_hex(m["primaire"]), "logo-couleur-marque.png")
    img.save(out / "logo-transparent.png")
    faits.append("logo-transparent.png")
    cote = max(img.size)
    carre = Image.new("RGBA", (cote, cote), (0, 0, 0, 0))
    carre.paste(img, ((cote - img.width) // 2, (cote - img.height) // 2), img)
    carre.save(out / "logo-carre.png")
    carre.resize((512, 512), Image.LANCZOS).save(out / "icone-512.png")
    carre.resize((64, 64), Image.LANCZOS).save(out / "favicon.ico", sizes=[(16, 16), (32, 32), (48, 48), (64, 64)])
    faits += ["logo-carre.png", "icone-512.png", "favicon.ico"]
    if nom:
        h = 200
        petit = img.resize((int(img.width * h / img.height), h), Image.LANCZOS)
        font = _police(96)
        tw = int(ImageDraw.Draw(Image.new("RGB", (10, 10))).textlength(nom, font=font))
        horiz = Image.new("RGBA", (petit.width + tw + 80, h + 40), (0, 0, 0, 0))
        horiz.paste(petit, (20, 20), petit)
        ImageDraw.Draw(horiz).text((petit.width + 50, 20 + (h - 96) // 2), nom, font=font, fill=_hex(m.get("primaire", "#111111")) + (255,))
        horiz.save(out / "logo-horizontal.png")
        faits.append("logo-horizontal.png")
    return f"[[image:{out / 'logo-carre.png'}]]\n{len(faits)} fichiers dans {out} : {', '.join(faits)}"


def business_cards(name: str, title: str, phone: str = "", email: str = "", website: str = "", output: str = "") -> str:
    """Cartes de visite (85 × 55 mm, recto et verso, couleurs de la charte) en PDF prêt à imprimer.

    Args:
        name: Nom.
        title: Fonction.
        phone: Téléphone.
        email: E-mail.
        website: Site.
        output: Fichier PDF (vide = workspace/cartes-de-visite.pdf).
    """
    from fpdf import FPDF

    m = _json(MARQUE, {})
    prim = _hex(m.get("primaire", "#1E3A8A"))
    sec = _hex(m.get("secondaire", "#F59E0B"))
    out = Path(output) if output else config.WORKSPACE / "cartes-de-visite.pdf"
    out.parent.mkdir(parents=True, exist_ok=True)
    pdf = FPDF(unit="mm", format=(85, 55))
    pdf.set_margins(6, 6, 6)
    pdf.set_auto_page_break(False)
    pdf.add_page()
    pdf.set_fill_color(*prim)
    pdf.rect(0, 0, 85, 8, "F")
    logo = m.get("logo", "")
    if logo and Path(logo).exists():
        pdf.image(logo, 62, 12, 17)
    pdf.set_xy(6, 14)
    pdf.set_font("Helvetica", "B", 15)
    pdf.set_text_color(*prim)
    pdf.cell(0, 8, name, new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 10)
    pdf.set_text_color(70, 70, 70)
    pdf.cell(0, 6, title, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)
    pdf.set_font("Helvetica", "", 9)
    for l in (phone, email, website):
        if l:
            pdf.cell(0, 5, l, new_x="LMARGIN", new_y="NEXT")
    pdf.set_fill_color(*sec)
    pdf.rect(0, 52, 85, 3, "F")
    pdf.add_page()
    pdf.set_fill_color(*prim)
    pdf.rect(0, 0, 85, 55, "F")
    if logo and Path(logo).exists():
        pdf.image(logo, 30, 13, 25)
    pdf.set_xy(0, 40)
    pdf.set_font("Helvetica", "B", 12)
    pdf.set_text_color(255, 255, 255)
    pdf.cell(85, 8, m.get("nom", name), align="C")
    pdf.output(str(out))
    return f"Cartes de visite créées (recto/verso) : {out}. Pour l'imprimeur : fond perdu 3 mm à ajouter si demandé."


def _blender():
    base = Path(r"C:\Program Files\Blender Foundation")
    candidats = [shutil.which("blender"), "/Applications/Blender.app/Contents/MacOS/Blender"]
    if base.exists():
        candidats += [str(d / "blender.exe") for d in sorted(base.glob("Blender*"), reverse=True)]
    return next((c for c in candidats if c and Path(c).exists()), None)


def blender_render(script: str = "", blend_file: str = "", output: str = "", samples: int = 32) -> str:
    """3D avec Blender en arrière-plan : exécute un script bpy (scène, objets, matériaux) et rend une image ; ou rend un fichier .blend existant.

    Args:
        script: Code Python bpy à exécuter (vide = scène de démonstration : sol, cube, sphère, lumière, caméra).
        blend_file: Fichier .blend à ouvrir avant le script (vide = scène neuve).
        output: Image PNG de sortie (vide = workspace/rendu.png).
        samples: Échantillons de rendu.
    """
    exe = _blender()
    if not exe:
        return "Blender est introuvable : installe-le (blender.org) ou ajoute-le au PATH."
    out = Path(output) if output else config.WORKSPACE / "rendu.png"
    out.parent.mkdir(parents=True, exist_ok=True)
    demo = ("import bpy\nbpy.ops.mesh.primitive_plane_add(size=20)\nbpy.ops.mesh.primitive_cube_add(location=(-1.5, 0, 1))\n"
            "bpy.ops.mesh.primitive_uv_sphere_add(location=(1.5, 0, 1))\nbpy.ops.object.shade_smooth()\n"
            "bpy.ops.object.light_add(type='SUN', location=(4, -4, 8))\nbpy.context.object.data.energy = 4\n"
            "bpy.ops.object.camera_add(location=(7, -7, 5), rotation=(1.1, 0, 0.785))\nbpy.context.scene.camera = bpy.context.object\n")
    fin = ("\nimport bpy\nsc = bpy.context.scene\n"
           "moteurs = [e.identifier for e in bpy.types.RenderSettings.bl_rna.properties['engine'].enum_items]\n"
           "sc.render.engine = 'BLENDER_EEVEE_NEXT' if 'BLENDER_EEVEE_NEXT' in moteurs else 'BLENDER_EEVEE'\n"
           f"sc.render.resolution_x, sc.render.resolution_y = 1280, 720\nsc.eevee.taa_render_samples = {int(samples)}\n"
           f"sc.render.filepath = r'{out}'\nbpy.ops.render.render(write_still=True)\n")
    sp = config.WORKSPACE / "blender_script.py"
    sp.write_text((script or demo) + fin, encoding="utf-8")
    args = [exe, "--background"] + ([str(Path(blend_file).expanduser())] if blend_file else []) + ["--python", str(sp)]
    r = _run(args, timeout=900)
    if out.exists():
        return f"[[image:{out}]]\nRendu Blender : {out}"
    return f"Pas d'image produite. Sortie Blender :\n{r}"


def design_tool(action: str = "canva_open", query: str = "", file_key: str = "", node_id: str = "") -> str:
    """Intégrations design : lire une maquette Figma (figma_read, via FIGMA_TOKEN) ou ouvrir Canva sur un modèle (canva_open) pour créer à l'écran.

    Args:
        action: "figma_read" ou "canva_open".
        query: Type de design pour Canva (ex. "affiche concert", "post instagram").
        file_key: Clé du fichier Figma (figma_read).
        node_id: Cadre Figma (facultatif).
    """
    if action == "figma_read":
        from outils_ia_plus import figma_to_code

        return figma_to_code(file_key, node_id)
    webbrowser.open(f"https://www.canva.com/templates/?query={query}" if query else "https://www.canva.com/")
    return "Canva ouvert sur les modèles : dis-moi lequel et je clique, je remplace les textes et j'exporte (see_screen)."


# --------------------------------------------------------------------------
# Création de contenu
# --------------------------------------------------------------------------

def content_ideas(niche: str, count: int = 10, format: str = "short", save: bool = True) -> str:
    """Idées, scripts et hooks de contenu pour une niche : rassemble tendances et sujets qui marchent, puis tu rédiges ; les idées sont gardées en mémoire.

    Args:
        niche: La niche (ex. "muscu débutants", "dev Python").
        count: Nombre d'idées.
        format: "short" (TikTok/Shorts), "long" (YouTube), "post" (Instagram/LinkedIn).
        save: Enregistrer les idées.
    """
    tendances = niche_trends(niche)[:1500]
    if save:
        base = _json(CONTENU, {})
        if not isinstance(base, dict):
            base = {}
        base.setdefault("idees", []).append({"niche": niche, "format": format, "t": datetime.now().isoformat(timespec="minutes"), "statut": "à rédiger"})
        _save(CONTENU, base)
    return (f"Tendances et matière pour « {niche} » :\n{tendances}\n\nProduis {count} idées au format {format}, chacune avec : titre, hook (première phrase, moins de 8 mots), "
            "angle, structure en 3 points, appel à l'action. Varie : tutoriel, erreur fréquente, avant/après, mythe, liste, histoire perso.")


def thumbnails(title: str, image_path: str = "", variants: int = 3, output_dir: str = "") -> str:
    """Génère plusieurs miniatures YouTube (1280×720) : gros titre lisible, dégradé, image de fond optionnelle, couleurs de la charte ; puis une planche comparative.

    Args:
        title: Texte de la miniature (court, 2 à 5 mots par ligne).
        image_path: Image de fond (photo, capture) facultative.
        variants: Nombre de variantes (mises en page différentes, 1 à 4).
        output_dir: Dossier de sortie (vide = workspace/miniatures).
    """
    from PIL import Image, ImageDraw

    m = _json(MARQUE, {})
    prim = _hex(m.get("primaire", "#FF3D00"))
    sec = _hex(m.get("secondaire", "#FFD600"))
    out = Path(output_dir) if output_dir else config.WORKSPACE / "miniatures"
    out.mkdir(parents=True, exist_ok=True)
    W, H = 1280, 720
    fond_src = None
    if image_path and Path(image_path).exists():
        fond_src = Image.open(image_path).convert("RGB")
        ratio = max(W / fond_src.width, H / fond_src.height)
        fond_src = fond_src.resize((int(fond_src.width * ratio), int(fond_src.height * ratio)), Image.LANCZOS)
        gx, gy = (fond_src.width - W) // 2, (fond_src.height - H) // 2
        fond_src = fond_src.crop((gx, gy, gx + W, gy + H))
    mots = title.upper().split()
    lignes = [" ".join(mots[i:i + 3]) for i in range(0, len(mots), 3)]
    fichiers = []
    layouts = [("gauche", (0.06, 0.55)), ("centre", (0.5, 0.5)), ("bas", (0.06, 0.78)), ("haut", (0.06, 0.2))]
    for i in range(max(1, min(int(variants), 4))):
        nom, (ax, ay) = layouts[i]
        img = fond_src.copy() if fond_src else Image.new("RGB", (W, H), prim if i % 2 == 0 else (20, 20, 30))
        if fond_src:
            img = Image.blend(img, Image.new("RGB", (W, H), (0, 0, 0)), 0.35)
        grad = Image.new("L", (W, H), 0)
        gd = ImageDraw.Draw(grad)
        for y in range(H):
            gd.line([(0, y), (W, y)], fill=int(200 * (y / H) ** 2))
        img.paste(Image.new("RGB", (W, H), (0, 0, 0)), (0, 0), grad)
        d = ImageDraw.Draw(img)
        taille = 118 if len(lignes) <= 2 else 92
        font = _police(taille)
        bloc_h = len(lignes) * (taille + 14)
        y0 = int(H * ay - bloc_h / 2)
        for k, l in enumerate(lignes):
            tw = d.textlength(l, font=font)
            x = int(W * ax) if nom != "centre" else int((W - tw) / 2)
            y = y0 + k * (taille + 14)
            for dx, dy in ((-4, -4), (4, -4), (-4, 4), (4, 4), (0, 6)):
                d.text((x + dx, y + dy), l, font=font, fill=(0, 0, 0))
            d.text((x, y), l, font=font, fill=(255, 255, 255) if k % 2 == 0 else sec)
        d.rectangle([0, H - 14, W, H], fill=sec)
        f = out / f"miniature-{i + 1}-{nom}.png"
        img.save(f)
        fichiers.append(f)
    lignes_planche = (len(fichiers) + 1) // 2
    planche = Image.new("RGB", (W + 30, (H // 2 + 20) * lignes_planche + 10), (40, 40, 40))
    for i, f in enumerate(fichiers):
        planche.paste(Image.open(f).resize((W // 2, H // 2)), (10 + (i % 2) * (W // 2 + 10), 10 + (i // 2) * (H // 2 + 20)))
    pf = out / "planche.png"
    planche.save(pf)
    return f"[[image:{pf}]]\n{len(fichiers)} miniatures dans {out}. Regarde la planche : laquelle se lit le mieux en petit ? (thumbnail_test pour le score)"


def thumbnail_test(folder: str = "") -> str:
    """Compare des miniatures : lisibilité à petite taille (contraste, encombrement) et recommande la meilleure.

    Args:
        folder: Dossier des miniatures (vide = workspace/miniatures).
    """
    from PIL import Image, ImageFilter, ImageStat

    d = Path(folder) if folder else config.WORKSPACE / "miniatures"
    fichiers = [f for f in d.glob("*.png") if not f.name.startswith("planche")]
    if not fichiers:
        return "Aucune miniature (thumbnails(titre) pour en créer)."
    scores = []
    for f in fichiers:
        img = Image.open(f).convert("L").resize((168, 94))
        contraste = ImageStat.Stat(img).stddev[0]
        encombrement = ImageStat.Stat(img.filter(ImageFilter.FIND_EDGES)).mean[0]
        score = contraste * 1.2 - max(0, encombrement - 25) * 1.5
        scores.append((score, f.name, contraste, encombrement))
    scores.sort(reverse=True)
    return ("\n".join(f"- {n} : score {s:.0f} (contraste {c:.0f}, encombrement {e:.0f})" for s, n, c, e in scores)
            + f"\nRecommandée : {scores[0][1]}. Idéal : contraste élevé, peu de détails fins, 3 à 5 mots.")


def platform_stats(platform: str = "youtube", days: int = 28) -> str:
    """Statistiques des plateformes : YouTube via l'API (YOUTUBE_KEY + YOUTUBE_CHANNEL_ID : abonnés, vues, dernières vidéos) ; TikTok / Instagram : ouvre le studio pour lire à l'écran.

    Args:
        platform: "youtube", "tiktok" ou "instagram".
        days: Fenêtre indicative.
    """
    p = platform.lower()
    if p != "youtube":
        pages = {"tiktok": "https://www.tiktok.com/tiktokstudio/analytics", "instagram": "https://business.facebook.com/latest/insights/overview"}
        webbrowser.open(pages.get(p, "https://www.google.com"))
        return f"Studio {platform} ouvert : lis les chiffres à l'écran (see_screen) et dis-les-moi, je les garde en mémoire."
    cle, chaine = getattr(config, "YOUTUBE_KEY", ""), getattr(config, "YOUTUBE_CHANNEL_ID", "")
    if not cle or not chaine:
        webbrowser.open("https://studio.youtube.com/")
        return "YOUTUBE_KEY / YOUTUBE_CHANNEL_ID manquent dans config.py : YouTube Studio ouvert, lis les chiffres à l'écran."
    import requests

    try:
        ch = requests.get("https://www.googleapis.com/youtube/v3/channels", params={"part": "statistics,contentDetails", "id": chaine, "key": cle}, timeout=15).json()["items"][0]
        st = ch["statistics"]
        uploads = ch["contentDetails"]["relatedPlaylists"]["uploads"]
        vids = requests.get("https://www.googleapis.com/youtube/v3/playlistItems", params={"part": "snippet", "playlistId": uploads, "maxResults": 10, "key": cle}, timeout=15).json()["items"]
        ids = ",".join(v["snippet"]["resourceId"]["videoId"] for v in vids)
        det = requests.get("https://www.googleapis.com/youtube/v3/videos", params={"part": "statistics,snippet", "id": ids, "key": cle}, timeout=15).json()["items"]
    except Exception as exc:  # noqa: BLE001
        return f"API YouTube : {exc}"
    base = _json(STATS, {})
    hist = base.setdefault("youtube", [])
    lignes = []
    for v in det:
        s, sn = v["statistics"], v["snippet"]
        lignes.append(f"- {sn['title'][:50]} ({sn['publishedAt'][:10]}) : {s.get('viewCount', 0)} vues, {s.get('likeCount', 0)} likes, {s.get('commentCount', 0)} commentaires")
        hist.append({"id": v["id"], "titre": sn["title"], "publie": sn["publishedAt"], "vues": int(s.get("viewCount", 0)), "t": datetime.now().isoformat(timespec="minutes")})
    base["youtube"] = hist[-500:]
    _save(STATS, base)
    return f"Chaîne : {st.get('subscriberCount')} abonnés, {st.get('viewCount')} vues, {st.get('videoCount')} vidéos.\nDernières vidéos :\n" + "\n".join(lignes)


def schedule_post(platform: str, path: str, title: str, when: str, description: str = "") -> str:
    """Publication programmée : enregistre la publication, puis à l'heure dite ouvre la page d'upload avec titre et description dans le presse-papiers et te prévient (l'upload se finit à l'écran).

    Args:
        platform: youtube, tiktok, instagram, linkedin.
        path: Fichier vidéo ou image.
        title: Titre.
        when: Date et heure "AAAA-MM-JJ HH:MM".
        description: Description / légende.
    """
    p = Path(path).expanduser()
    if not p.exists():
        return f"Introuvable : {p}"
    try:
        quand = datetime.strptime(when, "%Y-%m-%d %H:%M")
    except ValueError:
        return "Format attendu : AAAA-MM-JJ HH:MM."
    base = _json(CONTENU, {})
    if not isinstance(base, dict):
        base = {}
    base.setdefault("publications", []).append({"plateforme": platform, "fichier": str(p), "titre": title, "quand": when, "description": description, "statut": "programmée"})
    _save(CONTENU, base)
    pages = {"youtube": "https://studio.youtube.com/", "tiktok": "https://www.tiktok.com/tiktokstudio/upload", "instagram": "https://www.instagram.com/", "linkedin": "https://www.linkedin.com/feed/"}

    def alarme():
        delai = (quand - datetime.now()).total_seconds()
        if delai > 0:
            time.sleep(delai)
        try:
            import pyperclip

            pyperclip.copy(f"{title}\n\n{description}")
        except Exception:  # noqa: BLE001
            pass
        webbrowser.open(pages.get(platform, pages["youtube"]))
        try:
            import noyau

            noyau.hooks["dire"](f"C'est l'heure de publier « {title} » sur {platform} : la page est ouverte, le titre et la description sont dans le presse-papiers.")
        except Exception:  # noqa: BLE001
            pass

    threading.Thread(target=alarme, daemon=True).start()
    return f"Publication « {title} » programmée sur {platform} le {when} : je t'ouvrirai la page d'upload et je t'aiderai à cliquer (pas d'API d'upload sans connexion OAuth)."


def video_formats(video_path: str, formats: str = "9:16,1:1,16:9", output_dir: str = "") -> str:
    """Transforme une vidéo en plusieurs formats (9:16 vertical, 1:1 carré, 16:9, 4:5) par recadrage centré, prêts pour chaque plateforme.

    Args:
        video_path: Vidéo source.
        formats: Formats séparés par des virgules.
        output_dir: Dossier de sortie (vide = à côté de la source).
    """
    ff = _ffmpeg()
    if not ff:
        return "ffmpeg introuvable."
    p = Path(video_path).expanduser()
    if not p.exists():
        return f"Introuvable : {p}"
    out = Path(output_dir) if output_dir else p.parent
    out.mkdir(parents=True, exist_ok=True)
    faits = []
    for f in formats.split(","):
        f = f.strip()
        try:
            a, b = (int(x) for x in f.split(":"))
        except ValueError:
            continue
        filtre = f"crop='if(gt(iw/ih,{a}/{b}),ih*{a}/{b},iw)':'if(gt(iw/ih,{a}/{b}),ih,iw*{b}/{a})',scale='if(gt({a},{b}),1920,1080)':-2"
        dest = out / f"{p.stem}-{a}x{b}.mp4"
        _run([ff, "-y", "-i", str(p), "-vf", filtre, "-c:v", "libx264", "-preset", "fast", "-crf", "20", "-c:a", "aac", str(dest)], timeout=3600)
        if dest.exists():
            faits.append(dest.name)
    return f"{len(faits)} format(s) : {', '.join(faits)} dans {out}" if faits else "Aucun format produit."


def comment_ideas(video_id: str, count: int = 100) -> str:
    """Récupère les commentaires d'une vidéo YouTube (API) et les classe : questions, demandes de contenu, critiques, compliments ; en sort des idées de vidéos.

    Args:
        video_id: Id ou URL de la vidéo.
        count: Nombre de commentaires à lire.
    """
    cle = getattr(config, "YOUTUBE_KEY", "")
    m = re.search(r"(?:v=|youtu\.be/|shorts/)([\w-]{11})", video_id)
    vid = m.group(1) if m else video_id
    if not cle:
        webbrowser.open(f"https://www.youtube.com/watch?v={vid}")
        return "YOUTUBE_KEY manque : vidéo ouverte, fais défiler les commentaires et lis-les (see_screen)."
    import requests

    comms, token = [], None
    try:
        while len(comms) < count:
            params = {"part": "snippet", "videoId": vid, "maxResults": 100, "key": cle, "textFormat": "plainText"}
            if token:
                params["pageToken"] = token
            r = requests.get("https://www.googleapis.com/youtube/v3/commentThreads", params=params, timeout=15).json()
            comms += [i["snippet"]["topLevelComment"]["snippet"]["textDisplay"] for i in r.get("items", [])]
            token = r.get("nextPageToken")
            if not token:
                break
    except Exception as exc:  # noqa: BLE001
        return f"API YouTube : {exc}"
    cats: dict[str, list[str]] = {"questions": [], "demandes": [], "critiques": [], "compliments": [], "autres": []}
    for c in comms:
        b = c.lower()
        if "?" in c:
            cats["questions"].append(c)
        elif any(x in b for x in ("tu pourrais", "fais une vidéo", "prochaine vidéo", "tuto sur", "vidéo sur", "explique")):
            cats["demandes"].append(c)
        elif any(x in b for x in ("nul", "bof", "faux", "pas d'accord", "dommage", "trop long")):
            cats["critiques"].append(c)
        elif any(x in b for x in ("merci", "top", "génial", "super", "bravo", "excellent")):
            cats["compliments"].append(c)
        else:
            cats["autres"].append(c)
    base = _json(CONTENU, {})
    if not isinstance(base, dict):
        base = {}
    base.setdefault("commentaires", []).append({"video": vid, "t": datetime.now().isoformat(timespec="minutes"), "demandes": cats["demandes"][:30], "questions": cats["questions"][:30]})
    _save(CONTENU, base)
    return (f"{len(comms)} commentaires : " + ", ".join(f"{k} {len(v)}" for k, v in cats.items()) + "\n\nDemandes :\n" + "\n".join(f"- {c[:140]}" for c in cats["demandes"][:15])
            + "\n\nQuestions :\n" + "\n".join(f"- {c[:140]}" for c in cats["questions"][:15]) + "\n\nTransforme ça en 5 idées de vidéos classées par demande.")


def voice_over(text: str, output: str = "", voice: str = "", speed: float = 1.0) -> str:
    """Voix off générée en local (Kokoro) : écrit un fichier WAV à partir d'un texte, pour une vidéo.

    Args:
        text: Le texte à dire.
        output: Fichier WAV (vide = workspace/voix-off.wav).
        voice: Voix Kokoro (vide = celle de la config, ex. ff_siwis ; af_heart pour l'anglais).
        speed: Vitesse (1.0 normal).
    """
    import numpy as np
    import soundfile as sf

    try:
        from kokoro import KPipeline
    except Exception:  # noqa: BLE001
        return "Kokoro manque."
    v = voice or getattr(config, "KOKORO_VOICE", "ff_siwis")
    lang = v[0] if v and v[0] in "abefhijpz" else "f"
    pipe = KPipeline(lang_code=lang)
    morceaux = []
    for _g, _p, audio in pipe(text, voice=v, speed=float(speed), split_pattern=r"\n+|(?<=[.!?…])\s+"):
        morceaux.append(np.asarray(audio, dtype="float32"))
        morceaux.append(np.zeros(int(24000 * 0.25), dtype="float32"))
    if not morceaux:
        return "Rien généré."
    son = np.concatenate(morceaux)
    out = Path(output) if output else config.WORKSPACE / "voix-off.wav"
    out.parent.mkdir(parents=True, exist_ok=True)
    sf.write(str(out), son, 24000)
    return f"Voix off générée : {out} ({len(son)/24000:.1f} s, voix {v})."


def niche_trends(niche: str, region: str = "FR") -> str:
    """Tendances d'une niche : tendances Google du jour (flux RSS), actualités récentes et sujets qui montent sur le web.

    Args:
        niche: Le sujet.
        region: Pays (FR, US…).
    """
    import requests

    out = []
    try:
        rss = requests.get(f"https://trends.google.com/trending/rss?geo={region}", timeout=15, headers={"User-Agent": "Mozilla/5.0"}).text
        titres = re.findall(r"<title>(.*?)</title>", rss)[1:21]
        mots = niche.lower().split()
        lies = [t for t in titres if any(m in t.lower() for m in mots)]
        if lies:
            out.append(f"Tendances Google du jour liées à {niche} : " + ", ".join(lies))
        else:
            out.append(f"Tendances Google du jour (rien de directement lié à {niche}) : " + ", ".join(titres[:10]))
    except Exception:  # noqa: BLE001
        pass
    try:
        import tools

        out.append("Actualités et sujets :\n" + tools.web_search(f"{niche} tendance {datetime.now():%B %Y}")[:1500])
    except Exception as exc:  # noqa: BLE001
        out.append(f"Recherche impossible : {exc}")
    return "\n\n".join(out)


def comment_reply_drafts(comments: str, tone: str = "proche et sincère") -> str:
    """Brouillons de réponses aux commentaires (un par commentaire) que tu rédiges ensuite dans le ton du créateur.

    Args:
        comments: Les commentaires, un par ligne.
        tone: Ton voulu.
    """
    lignes = [l.strip("- ") for l in comments.splitlines() if l.strip()]
    if not lignes:
        return "Donne les commentaires."
    return (f"{len(lignes)} commentaires. Pour chacun, une réponse {tone} de 1 à 2 phrases, avec le prénom si visible, une question en retour quand c'est pertinent, jamais deux réponses identiques :\n"
            + "\n".join(f"{i}. {l}" for i, l in enumerate(lignes, start=1)))


def publish_reply(platform: str, text: str, url: str = "") -> str:
    """Publie une réponse à un commentaire : ouvre la page (ou le studio), met le texte dans le presse-papiers et t'indique où cliquer.

    Args:
        platform: youtube, tiktok, instagram.
        text: La réponse.
        url: Adresse de la vidéo ou du commentaire (vide = page des commentaires du studio).
    """
    import pyperclip

    pyperclip.copy(text)
    pages = {"youtube": "https://studio.youtube.com/", "tiktok": "https://www.tiktok.com/tiktokstudio/comments", "instagram": "https://www.instagram.com/direct/inbox/"}
    webbrowser.open(url or pages.get(platform, pages["youtube"]))
    return "Réponse dans le presse-papiers et page ouverte : clique sur « Répondre » sous le commentaire (see_screen), Ctrl+V, puis Publier. Je peux faire les clics."


def best_post_time(platform: str = "youtube") -> str:
    """Meilleure heure de publication : d'après tes propres statistiques (vues par jour et heure de publication), sinon les créneaux généralement les plus performants.

    Args:
        platform: youtube, tiktok, instagram, linkedin.
    """
    hist = _json(STATS, {}).get(platform, [])
    if len(hist) >= 5:
        vus: dict[tuple, list] = {}
        for v in hist:
            try:
                d = datetime.fromisoformat(v["publie"].replace("Z", "+00:00"))
            except Exception:  # noqa: BLE001
                continue
            vus.setdefault((d.strftime("%A"), d.hour), []).append(v.get("vues", 0))
        classement = sorted(((sum(x) / len(x), k) for k, x in vus.items()), reverse=True)[:5]
        return "D'après tes vidéos : " + " ; ".join(f"{j} {h}h ({m:.0f} vues en moyenne)" for m, (j, h) in classement)
    generiques = {"youtube": "jeudi et vendredi 17h à 19h, samedi 10h à 12h", "tiktok": "mardi 9h, jeudi 12h et 19h, samedi 11h",
                  "instagram": "mardi à jeudi 11h à 13h et 19h", "linkedin": "mardi à jeudi 8h à 10h"}
    return (f"Pas assez de données personnelles (platform_stats pour en collecter) ; créneaux généralement bons sur {platform} (heure de Paris) : "
            f"{generiques.get(platform, generiques['youtube'])}.")


def retention_analysis(video_id: str = "", csv_path: str = "") -> str:
    """Analyse de rétention : ouvre la courbe dans YouTube Studio (à lire à l'écran) ou analyse un export CSV de rétention (moments où les gens décrochent).

    Args:
        video_id: Id ou URL de la vidéo.
        csv_path: Export CSV « Durée de visionnage » de YouTube Studio.
    """
    if csv_path:
        import csv

        p = Path(csv_path).expanduser()
        if not p.exists():
            return f"Introuvable : {p}"
        pts = []
        with p.open(encoding="utf-8", errors="replace", newline="") as f:
            for row in csv.reader(f):
                try:
                    pts.append((float(row[0]), float(row[1].strip("%"))))
                except Exception:  # noqa: BLE001
                    continue
        if len(pts) < 5:
            return "CSV non reconnu (deux colonnes : position, % de spectateurs)."
        chutes = [(a[0], a[1] - b[1]) for a, b in zip(pts, pts[1:]) if a[1] - b[1] > 3]
        chutes.sort(key=lambda x: -x[1])
        return (f"Rétention : {pts[-1][1]:.0f} % à la fin. Plus grosses chutes : " + ", ".join(f"à {t:.0f} % de la vidéo (−{d:.0f} pts)" for t, d in chutes[:5])
                + ". Regarde ce qui se passe à ces moments (transition molle, digression, promesse non tenue).")
    m = re.search(r"(?:v=|youtu\.be/|shorts/)([\w-]{11})", video_id)
    vid = m.group(1) if m else video_id
    webbrowser.open(f"https://studio.youtube.com/video/{vid}/analytics/tab-overview/period-default" if vid else "https://studio.youtube.com/")
    return "Analytics ouvert : onglet « Engagement » > courbe de rétention. Regarde l'écran et dis-moi où ça chute (see_screen), ou exporte le CSV et donne-le-moi."


def hashtags_description(topic: str, platform: str = "youtube", keywords: str = "") -> str:
    """Description et hashtags pour une publication : mots-clés, chapitres, appel à l'action, hashtags adaptés à la plateforme.

    Args:
        topic: Sujet de la vidéo ou du post.
        platform: youtube, tiktok, instagram, linkedin.
        keywords: Mots-clés à inclure, séparés par des virgules.
    """
    limites = {"youtube": "3 à 5 hashtags, description 2 à 3 paragraphes avec mots-clés dans les 2 premières lignes, chapitres (00:00), liens",
               "tiktok": "3 à 6 hashtags dont 1 large et 2 de niche, légende de moins de 150 caractères",
               "instagram": "5 à 10 hashtags mélangeant gros et niche, légende avec hook en 1re ligne", "linkedin": "3 hashtags max, texte aéré, question finale"}
    mots = [k.strip() for k in keywords.split(",") if k.strip()] or topic.split()
    tags = ["#" + re.sub(r"[^a-z0-9]", "", m.lower()) for m in mots if len(m) > 2][:8]
    return (f"Sujet « {topic} », {platform} : {limites.get(platform, limites['youtube'])}.\nHashtags candidats : {' '.join(tags)}.\n"
            "Rédige la description finale (hook, contenu, appel à l'action, hashtags) et propose 2 titres alternatifs.")


# --------------------------------------------------------------------------
# Vidéo et audio
# --------------------------------------------------------------------------

def face_crop(video_path: str, output: str = "", ratio: str = "9:16") -> str:
    """Recadrage automatique qui suit le visage : détecte le visage sur toute la vidéo et recadre en vertical (9:16) centré dessus.

    Args:
        video_path: Vidéo source.
        output: Vidéo de sortie (vide = -vertical.mp4).
        ratio: "9:16" ou "1:1".
    """
    ff = _ffmpeg()
    if not ff:
        return "ffmpeg introuvable."
    try:
        import cv2
    except Exception:  # noqa: BLE001
        return "OpenCV manque."
    p = Path(video_path).expanduser()
    if not p.exists():
        return f"Introuvable : {p}"
    cap = cv2.VideoCapture(str(p))
    W, H = int(cap.get(cv2.CAP_PROP_FRAME_WIDTH)), int(cap.get(cv2.CAP_PROP_FRAME_HEIGHT))
    fps = cap.get(cv2.CAP_PROP_FPS) or 25
    casc = cv2.CascadeClassifier(cv2.data.haarcascades + "haarcascade_frontalface_default.xml")
    centres = []
    i = 0
    pas = max(1, int(fps))
    while True:
        ok, fr = cap.read()
        if not ok:
            break
        if i % pas == 0:
            faces = casc.detectMultiScale(cv2.cvtColor(fr, cv2.COLOR_BGR2GRAY), 1.2, 5, minSize=(60, 60))
            if len(faces):
                x, y, w, h = max(faces, key=lambda f: f[2] * f[3])
                centres.append((i / fps, x + w / 2))
        i += 1
    cap.release()
    a, b = (int(x) for x in ratio.split(":"))
    cw = int(H * a / b) if H * a / b <= W else W
    segments = []
    cx_prec = None
    for t, cx in centres:
        if cx_prec is None or abs(cx - cx_prec) > W * 0.08:
            segments.append((t, cx))
            cx_prec = cx
    if not segments:
        segments = [(0.0, W / 2)]
    expr = str(int(max(0, min(W - cw, segments[0][1] - cw / 2))))
    for t, cx in segments[1:]:
        x = int(max(0, min(W - cw, cx - cw / 2)))
        expr = f"if(gte(t,{t:.2f}),{x},{expr})"
    out = Path(output) if output else p.with_name(p.stem + "-vertical.mp4")
    r = _run([ff, "-y", "-i", str(p), "-vf", f"crop={cw}:{H}:'{expr}':0", "-c:v", "libx264", "-preset", "fast", "-crf", "20", "-c:a", "copy", str(out)], timeout=3600)
    return f"Vidéo recadrée sur le visage ({len(segments)} repositionnement(s)) : {out}" if out.exists() else f"Échec : {r}"


def royalty_free_music(mood: str, duration: int = 60, generate: bool = False) -> str:
    """Musique libre de droits adaptée à une ambiance : ouvre les bibliothèques gratuites (Pixabay, YouTube Audio Library, Free Music Archive) avec la bonne recherche, ou génère une nappe d'ambiance locale.

    Args:
        mood: Ambiance (calme, énergique, épique, lo-fi, corporate…).
        duration: Durée souhaitée en secondes.
        generate: True pour générer une nappe ambiante simple en local (WAV) au lieu de chercher.
    """
    if generate:
        import numpy as np
        import soundfile as sf

        sr = 44100
        t = np.linspace(0, duration, int(sr * duration), endpoint=False)
        base = {"calme": [110, 165, 220], "énergique": [130.8, 196, 261.6, 392], "épique": [65.4, 98, 130.8, 196], "lo-fi": [146.8, 220, 293.7], "corporate": [174.6, 261.6, 349.2]}
        freqs = next((v for k, v in base.items() if k in mood.lower()), [110, 165, 220])
        son = sum(np.sin(2 * np.pi * f * t) * (0.6 / (i + 1)) * (1 + 0.1 * np.sin(2 * np.pi * 0.1 * (i + 1) * t)) for i, f in enumerate(freqs))
        env = np.minimum(1, t / 3) * np.minimum(1, (duration - t) / 4)
        son = (son / np.max(np.abs(son)) * 0.4 * env).astype("float32")
        slug = re.sub(r"[^a-z]", "", mood.lower())
        out = config.WORKSPACE / f"ambiance-{slug}.wav"
        out.parent.mkdir(parents=True, exist_ok=True)
        sf.write(str(out), son, sr)
        return f"Nappe ambiante « {mood} » générée ({duration} s) : {out}"
    webbrowser.open(f"https://pixabay.com/fr/music/search/{mood}/")
    return (f"Pixabay Music ouvert sur « {mood} » (gratuit, sans attribution). Autres : YouTube Audio Library (studio.youtube.com > Bibliothèque audio), freemusicarchive.org. "
            "Dis-moi le titre qui te plaît et je le télécharge (download_file).")


def backup_rushes(source: str, destination: str, verify: bool = True) -> str:
    """Archive / sauvegarde des rushs vidéo : copie complète vers un autre disque avec vérification du nombre et de la taille des fichiers.

    Args:
        source: Dossier des rushs.
        destination: Dossier de sauvegarde.
        verify: Vérifier après copie.
    """
    s, d = Path(source).expanduser(), Path(destination).expanduser()
    if not s.is_dir():
        return f"Source introuvable : {s}"
    if IS_WINDOWS:
        subprocess.run(["robocopy", str(s), str(d), "/E", "/R:2", "/W:2", "/NP", "/NDL", "/NJH"], capture_output=True, creationflags=NO_WINDOW)
    else:
        subprocess.run(["rsync", "-a", str(s) + "/", str(d) + "/"], capture_output=True)
    if not verify:
        return "Copie lancée."

    def inv(p: Path):
        fs = [f for f in p.rglob("*") if f.is_file()]
        return len(fs), sum(f.stat().st_size for f in fs)

    ns, ts = inv(s)
    nd, td = inv(d)
    ok = nd >= ns and td >= ts
    return f"{'Sauvegarde vérifiée' if ok else 'ATTENTION, écart'} : source {ns} fichiers ({ts/1e9:.2f} Go), destination {nd} fichiers ({td/1e9:.2f} Go)."


def headphone_eq(preset: str = "flat", bands: str = "") -> str:
    """Égaliseur casque via Equalizer APO (s'il est installé) : préréglages (flat, bass, vocal, gaming, harman) ou bandes personnalisées.

    Args:
        preset: "flat", "bass", "vocal", "gaming", "harman" ou "custom".
        bands: Pour custom : "60:+3;250:-1;1000:0;4000:+2;12000:+1" (Hz:dB).
    """
    conf = Path(r"C:\Program Files\EqualizerAPO\config\config.txt")
    if not conf.exists():
        return "Equalizer APO n'est pas installé (sourceforge.net/projects/equalizerapo) : c'est lui qui applique l'égalisation au casque. Installe-le, choisis le casque, redémarre, puis réessaie."
    presets = {"flat": "GraphicEQ: 20 0; 20000 0", "bass": "GraphicEQ: 20 5; 60 5; 120 3; 250 1; 1000 0; 4000 0; 8000 1; 16000 1",
               "vocal": "GraphicEQ: 20 -2; 100 -2; 250 0; 1000 2; 2500 3; 4000 2; 8000 0; 16000 -1",
               "gaming": "GraphicEQ: 20 2; 100 2; 250 -1; 1000 0; 2500 2; 4000 4; 8000 3; 16000 1",
               "harman": "GraphicEQ: 20 4; 60 3; 120 1; 250 -1; 1000 0; 2000 2; 3000 3; 5000 1; 8000 -1; 16000 -3"}
    if preset == "custom" and bands:
        ligne = "GraphicEQ: " + "; ".join(f"{b.split(':')[0]} {b.split(':')[1].replace('+', '')}" for b in bands.split(";") if ":" in b)
    else:
        ligne = presets.get(preset, presets["flat"])
    try:
        conf.write_text(f"Preamp: -3 dB\n{ligne}\n", encoding="utf-8")
    except PermissionError:
        from outils_windows import run_elevated

        run_elevated(f"Set-Content -Path '{conf}' -Value \"Preamp: -3 dB`n{ligne}\"")
    return f"Égaliseur « {preset} » appliqué (Equalizer APO, immédiat)."


_SONS: dict = {"stream": None, "stop": threading.Event()}


def focus_sounds(kind: str = "pluie", minutes: int = 30, volume: float = 0.3, stop: bool = False) -> str:
    """Sons d'ambiance pour se concentrer, générés en local : pluie, bruit brun, rose, blanc, vent, feu ; s'arrêtent seuls.

    Args:
        kind: pluie, brun, rose, blanc, vent, feu.
        minutes: Durée.
        volume: 0 à 1.
        stop: True pour arrêter tout de suite.
    """
    import numpy as np
    import sounddevice as sd

    if stop or _SONS["stream"]:
        _SONS["stop"].set()
        if _SONS["stream"]:
            try:
                _SONS["stream"].stop()
                _SONS["stream"].close()
            except Exception:  # noqa: BLE001
                pass
            _SONS["stream"] = None
        if stop:
            return "Sons d'ambiance arrêtés."
    sr = 44100
    etat = {"brun": 0.0, "rose": np.zeros(7), "t": 0}
    rng = np.random.default_rng()
    noyau_goutte = np.exp(-np.linspace(0, 6, 200)).astype("float32")
    noyau_feu = np.exp(-np.linspace(0, 10, 80)).astype("float32")

    def cb(outdata, frames, t, status):
        blanc = rng.standard_normal(frames).astype("float32")
        if kind == "blanc":
            son = blanc * 0.15
        elif kind == "brun":
            acc = np.cumsum(blanc) * 0.02
            acc = acc - np.linspace(0, acc[-1], frames) + etat["brun"] * 0.99
            etat["brun"] = float(acc[-1])
            son = np.clip(acc, -1, 1) * 0.5
        elif kind in ("pluie", "vent", "feu"):
            acc = np.cumsum(blanc) * 0.02
            acc = acc - np.linspace(0, acc[-1], frames)
            son = np.clip(acc, -1, 1) * 0.35
            if kind == "pluie":
                gouttes = (rng.random(frames) > 0.9985).astype("float32") * rng.random(frames).astype("float32")
                son = son * 0.6 + np.convolve(gouttes, noyau_goutte, mode="same") * 0.5
            elif kind == "vent":
                mod = 0.5 + 0.5 * np.sin(2 * np.pi * 0.07 * (etat["t"] + np.arange(frames)) / sr)
                son = son * mod
            else:
                crepite = (rng.random(frames) > 0.9995).astype("float32")
                son = son * 0.5 + np.convolve(crepite, noyau_feu, mode="same") * 0.8
        else:  # rose (filtre de Paul Kellet)
            b = etat["rose"]
            out = np.empty(frames, dtype="float32")
            for i, w in enumerate(blanc):
                b[0] = 0.99886 * b[0] + w * 0.0555179
                b[1] = 0.99332 * b[1] + w * 0.0750759
                b[2] = 0.96900 * b[2] + w * 0.1538520
                b[3] = 0.86650 * b[3] + w * 0.3104856
                b[4] = 0.55000 * b[4] + w * 0.5329522
                b[5] = -0.7616 * b[5] - w * 0.0168980
                out[i] = (b[0] + b[1] + b[2] + b[3] + b[4] + b[5] + b[6] + w * 0.5362) * 0.11
                b[6] = w * 0.115926
            son = out
        etat["t"] += frames
        outdata[:, 0] = np.asarray(son, dtype="float32") * float(volume)

    _SONS["stop"].clear()
    _SONS["stream"] = sd.OutputStream(samplerate=sr, channels=1, dtype="float32", callback=cb)
    _SONS["stream"].start()

    def fin():
        if not _SONS["stop"].wait(minutes * 60):
            focus_sounds(stop=True)

    threading.Thread(target=fin, daemon=True).start()
    return f"Ambiance « {kind} » lancée pour {minutes} minutes (focus_sounds(stop=True) pour arrêter)."


def identify_song(seconds: int = 8) -> str:
    """Identifie la chanson qui passe (micro, comme Shazam).

    Args:
        seconds: Durée d'écoute.
    """
    import asyncio

    import sounddevice as sd
    import soundfile as sf

    try:
        from shazamio import Shazam
    except Exception:  # noqa: BLE001
        return "Le paquet shazamio manque."
    audio = sd.rec(int(seconds * 44100), samplerate=44100, channels=1, dtype="float32")
    sd.wait()
    f = config.WORKSPACE / "shazam.wav"
    f.parent.mkdir(parents=True, exist_ok=True)
    sf.write(str(f), audio, 44100)
    try:
        r = asyncio.run(Shazam().recognize(str(f)))
    except Exception as exc:  # noqa: BLE001
        return f"Reconnaissance impossible : {exc}"
    tr = r.get("track")
    if not tr:
        return "Aucune correspondance : rapproche le micro de la source ou réessaie plus longtemps."
    return f"C'est « {tr.get('title')} » de {tr.get('subtitle')}" + (f" ({tr.get('url')})" if tr.get("url") else "") + "."


STATIONS = {"france inter": "https://icecast.radiofrance.fr/franceinter-midfi.mp3", "france info": "https://icecast.radiofrance.fr/franceinfo-midfi.mp3",
            "fip": "https://icecast.radiofrance.fr/fip-midfi.mp3", "france culture": "https://icecast.radiofrance.fr/franceculture-midfi.mp3",
            "france musique": "https://icecast.radiofrance.fr/francemusique-midfi.mp3", "mouv": "https://icecast.radiofrance.fr/mouv-midfi.mp3",
            "rtl": "https://streaming.radio.rtl.fr/rtl-1-44-128", "europe 1": "https://europe1.lmn.fm/europe1.mp3", "rmc": "https://audio.bfmtv.com/rmcradio_128.mp3",
            "nrj": "https://scdn.nrjaudio.fm/adwz1/fr/30001/mp3_128.mp3", "skyrock": "https://icecast.skyrock.net/s/natio_mp3_128k",
            "nostalgie": "https://scdn.nrjaudio.fm/adwz1/fr/30601/mp3_128.mp3", "radio nova": "https://novazz.ice.infomaniak.ch/novazz-128.mp3",
            "rire et chansons": "https://scdn.nrjaudio.fm/adwz1/fr/30401/mp3_128.mp3"}


def radio_podcast(query: str, kind: str = "radio", stop: bool = False) -> str:
    """Radio ou podcast à la voix : lance une station (France Inter, FIP, RTL, NRJ, Skyrock, Nova…) ou cherche un podcast (annuaire Apple) et lance le dernier épisode.

    Args:
        query: Nom de la station ou du podcast.
        kind: "radio" ou "podcast".
        stop: True pour arrêter la lecture.
    """
    if stop:
        try:
            import tools

            return tools.media_control("stop")
        except Exception:  # noqa: BLE001
            return "Lecture arrêtée."
    q = query.lower().strip()
    if kind == "radio":
        url = next((u for k, u in STATIONS.items() if k in q or q in k), None)
        if not url:
            webbrowser.open(f"https://www.radio-browser.info/search?name={q}")
            return f"Station inconnue : annuaire ouvert pour « {query} », dis-moi laquelle et je la lance."
        webbrowser.open(url)
        return f"{query} en lecture dans le navigateur (radio_podcast(stop=True) ou « stop la radio » pour couper)."
    import requests

    try:
        res = requests.get("https://itunes.apple.com/search", params={"term": query, "media": "podcast", "country": "FR", "limit": 3}, timeout=15).json().get("results", [])
    except Exception as exc:  # noqa: BLE001
        return f"Recherche impossible : {exc}"
    if not res:
        return f"Aucun podcast pour « {query} »."
    p = res[0]
    ep = ""
    flux = p.get("feedUrl", "")
    if flux:
        try:
            xml = requests.get(flux, timeout=15).text
            m = re.search(r"<enclosure[^>]+url=\"([^\"]+)\"", xml)
            titre = re.search(r"<item>.*?<title>(.*?)</title>", xml, re.S)
            if m:
                webbrowser.open(m.group(1))
                ep = f" Dernier épisode lancé : « {titre.group(1).strip()[:80] if titre else '?'} »."
        except Exception:  # noqa: BLE001
            pass
    if not ep:
        webbrowser.open(p.get("collectionViewUrl", ""))
    return f"Podcast « {p.get('collectionName')} » ({p.get('artistName')}).{ep}"


TOOLS = [brand_kit, brand_document, logo_variants, business_cards, blender_render, design_tool, content_ideas, thumbnails,
         thumbnail_test, platform_stats, schedule_post, video_formats, comment_ideas, voice_over, niche_trends,
         comment_reply_drafts, publish_reply, best_post_time, retention_analysis, hashtags_description, face_crop,
         royalty_free_music, backup_rushes, headphone_eq, focus_sounds, identify_song, radio_podcast]
